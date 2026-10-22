"""Offline source snapshot acceptance; no model or Office application is started."""

from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
from types import SimpleNamespace
import warnings
import zipfile

from openpyxl import Workbook
import pytest

from agent_service.office_source import SourceError, snapshot_source


def task(tmp_path):
    root = tmp_path / "task"
    root.mkdir(parents=True)
    return root


def simple_docx(path, body, *, extra=None):
    content_types = b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/></Types>'
    document = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:body>' + body + '</w:body></w:document>').encode("utf-8")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("word/document.xml", document)
        for name, data in (extra or {}).items():
            z.writestr(name, data)


def simple_xlsx(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_csv_unicode_literals_and_numeric_preservation(tmp_path):
    source = tmp_path / "values.csv"
    raw = "项目,编号,金额,表达式\n纸,001,3.50,=1+2\n笔,2,-4,https://example.org\n".encode("utf-8-sig")
    source.write_bytes(raw)
    result = snapshot_source(source, task(tmp_path))
    assert result["source_id"] == "selected" and result["kind"] == "table"
    assert result["sha256"] == sha256(raw).hexdigest() and result["bytes"] == len(raw)
    assert result["observed"]["rows"] == [["纸", "001", 3.5, "=1+2"],
                                            ["笔", 2, -4, "https://example.org"]]
    assert source.read_bytes() == raw
    snapshot = tmp_path / "task" / "sources" / "source.csv"
    assert snapshot.read_bytes() == raw
    assert json.loads((snapshot.parent / "source.json").read_text(encoding="utf-8")) == result
    assert "path" not in result


def test_csv_precision_and_formula_like_text_remain_literal(tmp_path):
    source = tmp_path / "precise.csv"
    source.write_text("项目,值\n高精度,0.1234567890123456789\n负零,-0.00\n编号,0007\n公式,+SUM(1)\n",
                      encoding="utf-8")
    result = snapshot_source(source, task(tmp_path))
    assert result["observed"]["rows"] == [["高精度", "0.1234567890123456789"],
                                            ["负零", "-0.00"], ["编号", "0007"],
                                            ["公式", "+SUM(1)"]]


def test_xlsx_literal_values_and_exact_sheet_coverage(tmp_path):
    source = tmp_path / "table.xlsx"
    simple_xlsx(source, [["项目", "金额"], ["纸", 3], ["笔", 4.5]])
    result = snapshot_source(source, task(tmp_path))
    assert result["kind"] == "table"
    assert result["observed"] == {"headers": ["项目", "金额"],
                                  "rows": [["纸", 3], ["笔", 4.5]]}
    assert result["coverage"]["sheet"] == "数据"


def test_xlsx_multiple_sheets_rejected(tmp_path):
    source = tmp_path / "multiple.xlsx"
    workbook = Workbook()
    workbook.active.append(["项目", "金额"])
    workbook.active.append(["纸", 3])
    workbook.create_sheet("隐藏").sheet_state = "hidden"
    workbook.save(source)
    with pytest.raises(SourceError, match="workbook_sheets_unsupported"):
        snapshot_source(source, task(tmp_path))


@pytest.mark.parametrize("mutation", [
    lambda ws: ws.__setitem__("B2", "=SUM(1,2)"),
    lambda ws: setattr(ws.row_dimensions[2], "hidden", True),
    lambda ws: setattr(ws.column_dimensions["B"], "hidden", True),
    lambda ws: ws.merge_cells("A2:B2"),
    lambda ws: ws.__setitem__("B2", datetime(2026, 10, 17)),
])
def test_xlsx_ambiguous_or_active_cells_rejected(tmp_path, mutation):
    source = tmp_path / "bad.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "金额"])
    sheet.append(["纸", 3])
    mutation(sheet)
    workbook.save(source)
    with pytest.raises(SourceError):
        snapshot_source(source, task(tmp_path))


def test_docx_main_body_unicode_and_terminal_table(tmp_path):
    source = tmp_path / "notes.docx"
    body = ('<w:p><w:r><w:t>会议纪要</w:t></w:r></w:p>'
            '<w:p><w:r><w:t>预算共 7 元。</w:t></w:r></w:p>'
            '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>项目</w:t></w:r></w:p></w:tc>'
            '<w:tc><w:p><w:r><w:t>金额</w:t></w:r></w:p></w:tc></w:tr>'
            '<w:tr><w:tc><w:p><w:r><w:t>纸</w:t></w:r></w:p></w:tc>'
            '<w:tc><w:p><w:r><w:t>3</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')
    simple_docx(source, body)
    result = snapshot_source(source, task(tmp_path))
    assert result["kind"] == "document"
    assert result["observed"] == {"paragraphs": ["会议纪要", "预算共 7 元。"],
                                  "table": {"headers": ["项目", "金额"], "rows": [["纸", "3"]]}}
    assert result["coverage"]["scope"] == "main_body_only"


def test_docx_literal_hyphen_elements_in_paragraph_and_table(tmp_path):
    source = tmp_path / "hyphens.docx"
    body = ('<w:p><w:r><w:t>Amount</w:t><w:noBreakHyphen/>'
            '<w:t>due</w:t><w:softHyphen/><w:t>today</w:t></w:r></w:p>'
            '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Column</w:t></w:r></w:p></w:tc></w:tr>'
            '<w:tr><w:tc><w:p><w:r><w:t>one</w:t><w:noBreakHyphen/>'
            '<w:t>two</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')
    simple_docx(source, body)
    result = snapshot_source(source, task(tmp_path))
    assert result["observed"] == {
        "paragraphs": ["Amount\u2011due\u00adtoday"],
        "table": {"headers": ["Column"], "rows": [["one\u2011two"]]},
    }


@pytest.mark.parametrize("body", [
    '<w:p><w:r><w:t>Value</w:t><w:sym w:font="Wingdings" w:char="F0FC"/></w:r></w:p>',
    '<w:p><w:fldSimple w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>',
    '<w:p><w:r><w:pgNum/></w:r></w:p>',
    '<w:p><w:r><w:control/></w:r></w:p>',
    '<w:p><w:r><w:txbxContent><w:p><w:r><w:t>boxed</w:t></w:r></w:p>'
    '</w:txbxContent></w:r></w:p>',
    '<w:p><m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
    '<m:r><m:t>x</m:t></m:r></m:oMath></w:p>',
    '<w:p><w:r><w:t>first</w:t><w:br w:type="page"/><w:t>last</w:t></w:r></w:p>',
    '<w:p><w:hyperlink w:anchor="here"><w:r><w:t>link</w:t></w:r></w:hyperlink></w:p>',
    '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Title</w:t></w:r></w:p></w:tc></w:tr>'
    '<w:sdt><w:sdtContent><w:tr><w:tc><w:p><w:r><w:t>hidden row</w:t>'
    '</w:r></w:p></w:tc></w:tr></w:sdtContent></w:sdt></w:tbl>',
    '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Title</w:t></w:r></w:p></w:tc>'
    '<w:customXml><w:tc><w:p><w:r><w:t>hidden cell</w:t></w:r></w:p>'
    '</w:tc></w:customXml></w:tr></w:tbl>',
])
def test_docx_visible_unsupported_nodes_fail_instead_of_dropping_text(tmp_path, body):
    source = tmp_path / "unsupported-visible.docx"
    simple_docx(source, body)
    root = task(tmp_path)
    with pytest.raises(SourceError, match="document_unsupported_content"):
        snapshot_source(source, root)
    assert not (root / "sources" / "source.json").exists()


@pytest.mark.parametrize("body", [
    '<w:p><w:ins><w:r><w:t>修订</w:t></w:r></w:ins></w:p>',
    '<w:p><w:r><w:drawing/></w:r></w:p>',
    '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>字段</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
    '<w:p><w:r><w:t>表后内容</w:t></w:r></w:p>',
])
def test_docx_unsupported_content_rejected(tmp_path, body):
    source = tmp_path / "unsupported.docx"
    simple_docx(source, body)
    with pytest.raises(SourceError):
        snapshot_source(source, task(tmp_path))


@pytest.mark.parametrize("cell_body", [
    '<w:p><w:r><w:t>一</w:t></w:r></w:p><w:p><w:r><w:t>二</w:t></w:r></w:p>',
    '<w:p><w:r><w:t>一</w:t></w:r></w:p><w:tcPr><w:gridSpan w:val="2"/></w:tcPr>',
])
def test_docx_ambiguous_table_cell_rejected(tmp_path, cell_body):
    source = tmp_path / "ambiguous.docx"
    body = ('<w:tbl><w:tr><w:tc><w:p><w:r><w:t>字段</w:t></w:r></w:p></w:tc></w:tr>'
            '<w:tr><w:tc>' + cell_body + '</w:tc></w:tr></w:tbl>')
    simple_docx(source, body)
    with pytest.raises(SourceError):
        snapshot_source(source, task(tmp_path))


def test_external_relationship_macro_and_dtd_rejected(tmp_path):
    cases = [
        {"word/_rels/document.xml.rels": b'<Relationships><Relationship Target="https://x.example/" TargetMode="External"/></Relationships>'},
        {"word/vbaProject.bin": b"macro"},
        {"word/custom.xml": b'<!DOCTYPE x [<!ENTITY a "bad">]><x>&a;</x>'},
    ]
    for index, extra in enumerate(cases):
        source = tmp_path / f"danger-{index}.docx"
        simple_docx(source, '<w:p><w:r><w:t>文字</w:t></w:r></w:p>', extra=extra)
        root = tmp_path / f"task-{index}"
        root.mkdir()
        with pytest.raises(SourceError):
            snapshot_source(source, root)


def test_zip_traversal_and_duplicate_part_rejected(tmp_path):
    for index, added in enumerate(("../escape.xml", "word/document.xml")):
        source = tmp_path / f"archive-{index}.docx"
        with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", "<Types/>")
            z.writestr("word/document.xml", "<w:document/>")
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                z.writestr(added, "<x/>")
        root = tmp_path / f"archive-task-{index}"
        root.mkdir()
        with pytest.raises(SourceError, match="package_unsafe"):
            snapshot_source(source, root)


def test_compression_bomb_rejected_before_xml_parsing(tmp_path):
    source = tmp_path / "bomb.docx"
    simple_docx(source, '<w:p><w:r><w:t>ok</w:t></w:r></w:p>',
                 extra={"word/filler.xml": b"A" * 500_000})
    with pytest.raises(SourceError, match="package_unsafe"):
        snapshot_source(source, task(tmp_path))


def test_invalid_utf8_and_oversize_observation_rejected(tmp_path):
    invalid = tmp_path / "invalid.csv"
    invalid.write_bytes(b"a,b\n\xff,2\n")
    with pytest.raises(SourceError, match="csv_invalid"):
        snapshot_source(invalid, task(tmp_path))
    wide = tmp_path / "wide.csv"
    wide.write_text("a,b\n" + "\n".join(f"{i},{'x'*1900}" for i in range(20)) + "\n", encoding="utf-8")
    wide_root = tmp_path / "wide-task"
    wide_root.mkdir()
    with pytest.raises(SourceError, match="source_observation_oversize"):
        snapshot_source(wide, wide_root)


def test_writer_incompatible_controls_rejected_before_receipt(tmp_path):
    source = tmp_path / "multiline.csv"
    source.write_text('字段,说明\nA,"第一行\n第二行"\n', encoding="utf-8")
    root = task(tmp_path)
    with pytest.raises(SourceError, match="source_writer_unsupported"):
        snapshot_source(source, root)
    assert not (root / "sources" / "source.json").exists()


def test_malformed_archive_and_unsupported_format(tmp_path):
    source = tmp_path / "broken.docx"
    source.write_bytes(b"not a zip")
    with pytest.raises(SourceError, match="package_invalid"):
        snapshot_source(source, task(tmp_path))
    other = tmp_path / "script.ps1"
    other.write_text("Write-Host x", encoding="utf-8")
    with pytest.raises(SourceError, match="source_path_invalid"):
        snapshot_source(other, task(tmp_path / "another"))


def test_source_modified_during_copy_is_rejected(tmp_path, monkeypatch):
    import agent_service.office_source as module
    source = tmp_path / "mutable.csv"
    source.write_text("名字,金额\n笔,2\n", encoding="utf-8")
    root = task(tmp_path)
    original_stat = Path.stat

    def changed_stat(path, *args, **kwargs):
        value = original_stat(path, *args, **kwargs)
        if path == source and (root / "sources" / "source.csv").exists():
            return SimpleNamespace(st_dev=value.st_dev, st_ino=value.st_ino,
                                   st_size=value.st_size, st_mtime_ns=value.st_mtime_ns + 1,
                                   st_mode=value.st_mode)
        return value

    monkeypatch.setattr(module.Path, "stat", changed_stat)
    with pytest.raises(SourceError, match="source_changed"):
        snapshot_source(source, root)
    assert not (root / "sources" / "source.json").exists()


def test_snapshot_conflict_does_not_overwrite(tmp_path):
    source = tmp_path / "values.csv"
    source.write_text("项,数\nA,1\n", encoding="utf-8")
    root = task(tmp_path)
    first = snapshot_source(source, root)
    with pytest.raises(SourceError, match="source_snapshot_exists"):
        snapshot_source(source, root)
    assert json.loads((root / "sources" / "source.json").read_text(encoding="utf-8")) == first


def test_wps_bookmark_range_markers_do_not_drop_visible_text(tmp_path):
    source = tmp_path / "wps-bookmark.docx"
    simple_docx(source, '<w:p><w:pPr/><w:bookmarkStart w:id="0" w:name="_GoBack"/>'
                       '<w:r><w:rPr/><w:t>资料归档</w:t></w:r><w:bookmarkEnd w:id="0"/></w:p>')
    result = snapshot_source(source, task(tmp_path))
    assert result["observed"]["paragraphs"] == ["资料归档"]


def test_wps_row_property_exceptions_are_nonvisible_formatting(tmp_path):
    source = tmp_path / "wps-table.docx"
    simple_docx(source, '<w:p><w:r><w:t>标题</w:t></w:r></w:p><w:tbl><w:tblPr/><w:tblGrid/>'
                       '<w:tr><w:tblPrEx><w:tblBorders/></w:tblPrEx><w:tc><w:tcPr/>'
                       '<w:p><w:r><w:t>字段</w:t></w:r></w:p></w:tc></w:tr>'
                       '<w:tr><w:tblPrEx/><w:tc><w:p><w:r><w:t>原值</w:t></w:r></w:p></w:tc></w:tr>'
                       '</w:tbl>')
    result = snapshot_source(source, task(tmp_path))
    assert result["observed"]["table"] == {"headers": ["字段"], "rows": [["原值"]]}


def test_selected_source_link_or_task_source_rejected(tmp_path):
    source = tmp_path / "original.csv"
    source.write_text("项,数\nA,1\n", encoding="utf-8")
    root = task(tmp_path)
    internal = root / "internal.csv"
    internal.write_bytes(source.read_bytes())
    with pytest.raises(SourceError, match="source_path_invalid"):
        snapshot_source(internal, root)
    link = tmp_path / "linked.csv"
    try:
        link.symlink_to(source)
    except OSError:
        pytest.skip("Symlinks unavailable on this host")
    with pytest.raises(SourceError, match="source_path_invalid"):
        snapshot_source(link, root)
