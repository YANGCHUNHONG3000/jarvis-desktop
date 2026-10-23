"""Offline gates for the limited supervised edit adapter (no native Office)."""
from openpyxl import Workbook
import pytest
import zipfile

from agent_service.computer import ComputerSession, read_record, write_new
from agent_service.computer_interaction import prepare
from agent_service.computer_goals import _save_goal
from agent_service import office_edit_adapter as adapter


def setup_plan(tmp_path, *, both=False):
    session = ComputerSession.create(tmp_path / "sessions")
    source = tmp_path / "input.xlsx"
    book = Workbook()
    book.active["A1"] = 2
    book.active["A2"] = "=SUM(A1:A1)"
    book.save(source)
    book.close()
    edits = [{"path": str(source), "changes": [
        {"sheet": "Sheet", "cell": "A1", "before": 2, "after": 3}]}]
    if both:
        from agent_service.office_edit import W
        word = tmp_path / "input.docx"
        with zipfile.ZipFile(word, "w", compression=zipfile.ZIP_DEFLATED) as out:
            out.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
            out.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
            out.writestr("word/document.xml", f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>before</w:t></w:r></w:p></w:body></w:document>')
        edits.append({"path": str(word), "changes": [
            {"paragraph_index": 0, "before": "before", "after": "after"}]})
    plan = prepare(session, "编辑 Word 和 Excel" if both else "把 A1 改为 3", edits=edits)
    return session, plan


def test_offline_review_and_immutable_identity(tmp_path):
    session, plan = setup_plan(tmp_path)
    store = tmp_path / "adapters"
    first = adapter.ensure_candidate(store, session, plan, proposer="author")
    candidate = first["candidate"]
    assert adapter.ensure_candidate(store, session, plan, proposer="author") == first
    check = adapter.offline(store, candidate)
    assert check["passed"] and [v["kind"] for v in check["fixtures"]] == ["document", "spreadsheet"]
    assert adapter.offline(store, candidate) == check
    with pytest.raises(adapter.AdapterError, match="review_invalid"):
        adapter.record_review(store, candidate, reviewer="author", note="looks fine")
    review = adapter.record_review(store, candidate, reviewer="independent", note="Reviewed fixture and contract")
    assert review == adapter.record_review(store, candidate, reviewer="independent",
                                           note="Reviewed fixture and contract")
    assert adapter.status(store, candidate)["reviewed"] is True
    with pytest.raises(adapter.AdapterError, match="trial_required"):
        adapter.activate(store, candidate, runs=1)


def test_changed_implementation_invalidates_candidate(tmp_path, monkeypatch):
    session, plan = setup_plan(tmp_path)
    store = tmp_path / "adapters"
    candidate = adapter.ensure_candidate(store, session, plan, proposer="author")["candidate"]
    monkeypatch.setattr(adapter, "implementation_sha256", lambda: "0" * 64)
    with pytest.raises(adapter.AdapterError, match="edit_adapter_changed"):
        adapter.status(store, candidate)


def test_fake_native_success_cannot_pass_trial(tmp_path, monkeypatch):
    session, plan = setup_plan(tmp_path, both=True)
    store = tmp_path / "adapters"
    candidate = adapter.ensure_candidate(store, session, plan, proposer="author")["candidate"]
    adapter.offline(store, candidate)
    adapter.record_review(store, candidate, reviewer="independent", note="reviewed")
    monkeypatch.setattr(adapter, "_run", lambda *args: {"ok": True, "state": "verified", "artifacts": []})
    with pytest.raises(Exception):
        adapter.trial(store, candidate, session, plan, approval=plan["sha256"])
    root = store / candidate
    assert (root / "trial-intent.json").is_file()
    assert not (root / "trial-proof.json").exists()
    with pytest.raises(adapter.AdapterError, match="trial_required"):
        adapter.activate(store, candidate)


def test_refresh_plan_keeps_exact_goal_and_sources(tmp_path):
    session, plan = setup_plan(tmp_path)
    fresh = adapter.refresh_plan(session, plan)
    assert fresh["goal"] == plan["goal"]
    assert fresh["instruction"] == plan["instruction"]
    assert fresh["sha256"] != plan["sha256"]


def test_adapter_owns_active_recovery_without_another_quota_slot(tmp_path):
    session, plan = setup_plan(tmp_path)
    store = tmp_path / "adapters"
    candidate = adapter.ensure_candidate(store, session, plan, proposer="author")["candidate"]
    root = store / candidate
    write_new(root / "active-1-intent.json", {"candidate": candidate,
              "plan_sha256": plan["sha256"], "session_root": str(session.root)})
    child = _save_goal(session, {"schema": 3, "mode": "agent", "instruction": plan["instruction"],
             "goal": plan["goal"], "max_actions": plan["max_actions"],
             "max_seconds": plan["max_seconds"], "prepared_at": plan["prepared_at"] + 1,
             "recovery": {"parent_run": str(session.root / "goals" / (plan["sha256"] + ".run"))}})
    assert adapter.owns_recovery(store, session, child) == {
        "candidate": candidate, "phase": "active", "index": 1}
    assert adapter._resume_parent(root, session, "active-1", read_record(
        root / "active-1-intent.json")) == str(session.root / "goals" / (plan["sha256"] + ".run"))


def reviewed(tmp_path, *, both=False):
    session, plan = setup_plan(tmp_path, both=both)
    store = tmp_path / "adapters"
    candidate = adapter.ensure_candidate(store, session, plan, proposer="author")["candidate"]
    adapter.offline(store, candidate)
    adapter.record_review(store, candidate, reviewer="reviewer", note="Checked source and preservation")
    return session, plan, store, candidate


def saved_fake_outcome(session, plan, *, ok, state):
    """A durable synthetic agent outcome for bookkeeping tests only."""
    run = session.root / "goals" / (plan["sha256"] + ".run")
    run.mkdir()
    outcome = {"ok": ok, "state": state, "run": str(run)}
    write_new(run / "outcome.json", outcome)
    return outcome


def test_expired_preview_does_not_write_trial_intent(tmp_path, monkeypatch):
    session, plan, store, candidate = reviewed(tmp_path)
    monkeypatch.setattr(adapter.time, "time", lambda: plan["prepared_at"] + 121)
    with pytest.raises(adapter.AdapterError, match="preview_expired"):
        adapter.trial(store, candidate, session, plan, approval=plan["sha256"])
    assert not (store / candidate / "trial-intent.json").exists()


def test_changed_source_does_not_write_trial_intent(tmp_path):
    session, plan, store, candidate = reviewed(tmp_path)
    original = tmp_path / "input.xlsx"
    original.write_bytes(original.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="edit_source_changed"):
        adapter.trial(store, candidate, session, plan, approval=plan["sha256"])
    assert not (store / candidate / "trial-intent.json").exists()


def test_unknown_trial_is_not_replayed_or_activated(tmp_path, monkeypatch):
    session, plan, store, candidate = reviewed(tmp_path, both=True)
    monkeypatch.setattr(adapter, "_run", lambda *args: saved_fake_outcome(
        session, plan, ok=False, state="unknown"))
    result = adapter.trial(store, candidate, session, plan, approval=plan["sha256"])
    assert result["proof"] is None
    assert read_record(store / candidate / "trial-result.json")["outcome"]["state"] == "unknown"
    with pytest.raises(adapter.AdapterError, match="plan_already_started"):
        adapter.trial(store, candidate, session, plan, approval=plan["sha256"])
    with pytest.raises(adapter.AdapterError, match="trial_required"):
        adapter.activate(store, candidate)


def test_single_kind_cannot_start_shared_profile_trial(tmp_path):
    session, plan, store, candidate = reviewed(tmp_path)
    with pytest.raises(adapter.AdapterError, match="trial_requires_both_kinds"):
        adapter.trial(store, candidate, session, plan, approval=plan["sha256"])
    assert not (store / candidate / "trial-intent.json").exists()


def test_total_changes_limit_applies_across_files(tmp_path):
    session = ComputerSession.create(tmp_path / "sessions")
    edits = []
    for index in (1, 2):
        source = tmp_path / f"source-{index}.xlsx"
        book = Workbook()
        for row in range(1, 12):
            book.active.cell(row, 1, row)
        book.save(source)
        book.close()
        edits.append({"path": str(source), "changes": [
            {"sheet": "Sheet", "cell": f"A{row}", "before": row, "after": row + 1}
            for row in range(1, 12)]})
    plan = prepare(session, "编辑两份工作簿", edits=edits)
    with pytest.raises(adapter.AdapterError, match="edit_adapter_change_limit"):
        adapter.ensure_candidate(tmp_path / "adapters", session, plan, proposer="author")
    assert not (tmp_path / "adapters").exists()


def test_foreign_recovery_cannot_be_stamped_as_adapter_run(tmp_path):
    session, plan, store, candidate = reviewed(tmp_path)
    foreign = _save_goal(session, {"schema": 3, "mode": "agent", "instruction": plan["instruction"],
        "goal": plan["goal"], "max_actions": plan["max_actions"], "max_seconds": plan["max_seconds"],
        "prepared_at": plan["prepared_at"],
        "recovery": {"parent_run": str(session.root / "goals" / (plan["sha256"] + ".run"))}})
    assert adapter.owns_recovery(store, session, foreign) is None
    with pytest.raises(adapter.AdapterError, match="recovery_unowned"):
        adapter.resume(store, candidate, session, foreign, approval=foreign["sha256"])
    assert not list((store / candidate).glob("*-resume-*-intent.json"))


def test_active_pause_and_recovery_use_one_consumed_slot(tmp_path, monkeypatch):
    """Synthetic outcomes test bookkeeping; production proof uses _verified."""
    session, plan, store, candidate = reviewed(tmp_path)
    root = store / candidate
    write_new(root / "trial-intent.json", {"candidate": candidate,
        "plan_sha256": "f" * 64, "session_root": str(session.root)})
    write_new(root / "activation-1.json", {"fixture": True})
    monkeypatch.setattr(adapter, "_activation_limit", lambda _: 1)
    monkeypatch.setattr(adapter, "_run", lambda *args: saved_fake_outcome(
        session, plan, ok=False, state="partial"))
    first = adapter.execute(store, candidate, session, plan, approval=plan["sha256"])
    assert first["proof"] is None
    assert len(adapter._active_records(root)) == 1
    child = _save_goal(session, {"schema": 3, "mode": "agent", "instruction": plan["instruction"],
        "goal": plan["goal"], "max_actions": plan["max_actions"], "max_seconds": plan["max_seconds"],
        "prepared_at": plan["prepared_at"],
        "recovery": {"parent_run": str(session.root / "goals" / (plan["sha256"] + ".run"))}})
    monkeypatch.setattr(adapter, "_run", lambda *args: saved_fake_outcome(
        session, child, ok=True, state="verified"))
    monkeypatch.setattr(adapter, "_verified", lambda *args: {"plan_sha256": child["sha256"],
        "run": str(session.root / "goals" / (child["sha256"] + ".run")),
        "artifacts": [], "artifacts_sha256": "0" * 64,
        "evidence_sha256": "1" * 64})
    recovered = adapter.resume(store, candidate, session, child, approval=child["sha256"])
    assert recovered["proof"] is not None
    assert (root / "active-1-proof.json").exists()
    monkeypatch.setattr(adapter, "_proof", lambda _root, path, _intent: read_record(path))
    assert len(adapter._active_records(root)) == 1


def test_activation_is_idempotent_and_total_quota(tmp_path, monkeypatch):
    """Synthetic saved proofs exercise counters; real proof verification is separate."""
    session, plan, store, candidate = reviewed(tmp_path)
    root = store / candidate
    write_new(root / "trial-intent.json", {"candidate": candidate,
        "plan_sha256": plan["sha256"], "session_root": str(session.root)})
    trial_proof = {"plan_sha256": plan["sha256"], "run": "fixture",
                   "kinds": ["document", "spreadsheet"]}
    write_new(root / "trial-proof.json", trial_proof)
    trial_outcome = saved_fake_outcome(session, plan, ok=True, state="verified")
    write_new(root / "trial-result.json", {"plan_sha256": plan["sha256"],
        "outcome": trial_outcome, "proof": trial_proof})
    monkeypatch.setattr(adapter, "_proof", lambda _root, path, _intent: read_record(path))
    one = adapter.activate(store, candidate, runs=1)
    activation_bytes = (root / "activation-1.json").read_bytes()
    assert one["enabled_total"] == 1 and one["remaining"] == 1
    assert adapter.activate(store, candidate, runs=1) == one
    assert (root / "activation-1.json").read_bytes() == activation_bytes
    write_new(root / "active-1-intent.json", {"candidate": candidate,
        "plan_sha256": "e" * 64, "session_root": str(session.root)})
    active_proof = {"plan_sha256": "e" * 64, "run": "fixture"}
    write_new(root / "active-1-proof.json", active_proof)
    fake_active = {"sha256": "e" * 64}
    active_outcome = saved_fake_outcome(session, fake_active, ok=True, state="verified")
    write_new(root / "active-1-result.json", {"outcome": active_outcome, "proof": active_proof})
    three = adapter.activate(store, candidate, runs=3)
    assert three["enabled_total"] == 3 and three["used"] == 1 and three["remaining"] == 2
    assert adapter.activate(store, candidate, runs=3) == three


def test_orphan_active_result_and_changed_outcome_are_rejected(tmp_path):
    session, plan, store, candidate = reviewed(tmp_path)
    root = store / candidate
    write_new(root / "active-2-result.json", {"outcome": {"ok": True}})
    with pytest.raises(adapter.AdapterError, match="active_records_changed"):
        adapter.status(store, candidate)
    (root / "active-2-result.json").unlink()
    write_new(root / "trial-intent.json", {"candidate": candidate,
        "plan_sha256": plan["sha256"], "session_root": str(session.root)})
    saved_fake_outcome(session, plan, ok=False, state="partial")
    write_new(root / "trial-result.json", {"plan_sha256": plan["sha256"],
        "outcome": {"ok": False, "state": "unknown", "run": "changed"}, "proof": None})
    with pytest.raises(adapter.AdapterError, match="records_changed"):
        adapter.status(store, candidate)


def test_orphan_resume_records_cannot_free_trial_or_active_slot(tmp_path):
    session, plan, store, candidate = reviewed(tmp_path, both=True)
    root = store / candidate
    write_new(root / "trial-resume-1-intent.json", {"parent_run": "missing",
        "plan_sha256": "e" * 64})
    with pytest.raises(adapter.AdapterError, match="records_changed"):
        adapter.status(store, candidate)
    with pytest.raises(adapter.AdapterError, match="records_changed"):
        adapter.trial(store, candidate, session, plan, approval=plan["sha256"])
    (root / "trial-resume-1-intent.json").unlink()
    write_new(root / "active-1-resume-1-intent.json", {"parent_run": "missing",
        "plan_sha256": "e" * 64})
    with pytest.raises(adapter.AdapterError, match="active_records_changed"):
        adapter.status(store, candidate)


def test_missing_resume_child_with_claim_is_not_skipped(tmp_path):
    session, plan, store, candidate = reviewed(tmp_path)
    root = store / candidate
    write_new(root / "trial-intent.json", {"candidate": candidate,
        "plan_sha256": plan["sha256"], "session_root": str(session.root)})
    parent = session.root / "goals" / (plan["sha256"] + ".run")
    parent.mkdir()
    child_sha = "e" * 64
    write_new(root / "trial-resume-1-intent.json", {"parent_run": str(parent),
        "plan_sha256": child_sha})
    write_new(parent / "continuation-claim.json", {"child": child_sha})
    with pytest.raises(adapter.AdapterError, match="recovery_changed"):
        adapter._resume_parent(root, session, "trial", read_record(root / "trial-intent.json"))


def test_reconcile_only_finishes_saved_terminal_outcome(tmp_path):
    session, plan, store, candidate = reviewed(tmp_path)
    root = store / candidate
    write_new(root / "trial-intent.json", {"candidate": candidate,
        "plan_sha256": plan["sha256"], "session_root": str(session.root)})
    outcome = saved_fake_outcome(session, plan, ok=False, state="partial")
    first = adapter.reconcile(store, candidate)
    assert first["records_finalized"] == 1
    assert read_record(root / "trial-result.json") == {
        "plan_sha256": plan["sha256"], "outcome": outcome, "proof": None}
    assert adapter.reconcile(store, candidate)["records_finalized"] == 0
    assert not (root / "trial-proof.json").exists()
