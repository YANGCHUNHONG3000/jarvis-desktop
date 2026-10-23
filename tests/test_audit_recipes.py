"""Fixed complete audit recipes and edited-workbook provenance gates."""

from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import tempfile
import uuid

from openpyxl import Workbook
import pytest

from agent_service import (adapter_development as development, artifact_source, assistant,
                           gap_recovery, local_skills as skills, task_history, workbook_source,
                           workbook_source_host)
from agent_service.edited_artifact_source import prepare_reference
from agent_service.verified_edit_document import snapshot_reference as snapshot_document


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "data" / "assistant_tasks"
EDITED = BASE / "task-5d651a699ec54962a43472883c3da374" / "edit" / "edited-spreadsheet.xlsx"
COMPLEX_WORD = BASE / "task-1d52d21e59a241e693601607106075fc" / "edit" / "edited-document.docx"
REAL_WORKBOOK_GAP = BASE / "task-2c93d9390d324c3db90ce661ddd391a8"


def _manifest(recipe):
    return {"schema": 1, "id": "audit-complete", "version": 1,
            "name": "完整审计报告", "recipe": recipe, "sum_columns": [], "draft": False}


def _native(path, directory, _stop, _deadline):
    observed, _ = workbook_source._parse(path)
    workbook_source.office_source._write_json_new(
        directory / "native-request.json", {"sha256": workbook_source._digest(path)})
    result = {"version": 1, "source_sha256": workbook_source._digest(path),
              "sheets": observed["sheets"], "application_verified": True,
              "file_unchanged": True,
              "app": {"progid": "KET.Application", "name": "unit-test", "version": "test"},
              "calculation": "owned_formula_ranges_in_dependency_order"}
    workbook_source.office_source._write_json_new(directory / "source-native.json", result)
    return result


def test_workbook_audit_includes_every_cell_and_rejects_overflow():
    candidate = _manifest("workbook_audit_report")
    source = skills._fixture(candidate)
    steps = skills.build_steps(candidate, source)
    assert len(steps) == 2
    sheet = steps[0]["data"]
    assert sheet["sum_columns"] == [] and len(sheet["rows"]) == 4
    assert sheet["rows"] == [
        ["甲", "visible", "A1", "数字", 2, "", 2],
        ["甲", "visible", "B1", "公式", None, "=A1*2", 4],
        ["乙", "hidden", "A1", "空白", None, "", None],
        ["乙", "hidden", "B1", "文本", "=A1", "", "=A1"],
    ]
    assert steps[1]["action"] == "document_from_sheet" and steps[1]["from"] == "s1"
    assert "不保留可执行公式" in steps[1]["paragraphs"][0]
    oversized = deepcopy(source)
    oversized["observed"]["sheets"] = [{"name": "大页", "range": "A1:A201", "state": "visible",
                                         "values": [[i] for i in range(201)], "formulas": []}]
    oversized["coverage"]["sheet_count"] = 1
    with pytest.raises(skills.SkillError, match="skill_audit_output_row_limit"):
        skills.build_steps(candidate, oversized)
    wrong = deepcopy(source)
    wrong["observed"]["sheets"][0]["formulas"][0]["result"] = 999
    with pytest.raises(ValueError):
        skills.build_steps(candidate, wrong)


def test_document_structure_audit_retains_order_and_metadata():
    candidate = _manifest("document_structure_report")
    source = skills._fixture(candidate)
    steps = skills.build_steps(candidate, source)
    assert len(steps[0]["data"]["rows"]) == len(source["observed"]["records"])
    assert [row[0] for row in steps[0]["data"]["rows"]] == [
        record["position"] for record in source["observed"]["records"]]
    assert '"table_index":0' in steps[0]["data"]["rows"][1][2]
    assert steps[1]["from"] == "s1" and "不重建" in steps[1]["paragraphs"][0]


def test_audit_recipes_reject_malformed_or_legacy_observations():
    structured = _manifest("document_structure_report")
    source = skills._fixture(structured)
    malformed = deepcopy(source)
    malformed["observed"]["records"][0].pop("position")
    with pytest.raises(ValueError, match="structured_document_observation_invalid"):
        skills.build_steps(structured, malformed)
    malformed = deepcopy(source)
    malformed["observed"]["records"].append({"position": "body/9", "type": "unknown",
                                               "structure": {}, "text": "stray"})
    with pytest.raises(ValueError, match="structured_document_observation_invalid"):
        skills.build_steps(structured, malformed)
    legacy = deepcopy(source)
    legacy["kind"] = "document"
    with pytest.raises(skills.SkillError, match="skill_source_kind_mismatch"):
        skills.build_steps(structured, legacy)
    workbook = _manifest("workbook_audit_report")
    with pytest.raises(skills.SkillError, match="skill_source_kind_mismatch"):
        skills.build_steps(workbook, legacy)


def test_preserved_complex_word_can_feed_complete_fixed_audit(tmp_path):
    if not COMPLEX_WORD.is_file():
        pytest.skip("preserved complex Word edit unavailable")
    ref = prepare_reference(COMPLEX_WORD, sha256(COMPLEX_WORD.read_bytes()).hexdigest(), BASE)
    task = tmp_path / "word-task"
    task.mkdir()
    snapshot = snapshot_document(ref, task, BASE)
    assert snapshot["kind"] == "document_structure"
    steps = skills.build_steps(_manifest("document_structure_report"), snapshot)
    assert len(steps[0]["data"]["rows"]) == len(snapshot["observed"]["records"])
    assert len(steps[0]["data"]["rows"]) > 2


def test_offline_gate_covers_both_audit_recipes(tmp_path):
    with tempfile.TemporaryDirectory(prefix="ar-") as short:
        for index, recipe in enumerate(("workbook_audit_report", "document_structure_report")):
            candidate = _manifest(recipe)
            store = Path(short) / str(index)
            digest = skills.register_candidate(store, candidate)["sha256"]
            record = skills.verify_offline(store, candidate["id"], digest)
            assert record["passed"] and record["kind"] == "offline_contract"


def test_real_edited_workbook_uses_native_complete_snapshot_and_chain(tmp_path, monkeypatch):
    if not EDITED.is_file():
        pytest.skip("preserved edited workbook unavailable")
    monkeypatch.setattr(workbook_source_host, "native_snapshot_supervised", _native)
    ref = prepare_reference(EDITED, sha256(EDITED.read_bytes()).hexdigest(), BASE)
    task = tmp_path / "task"
    task.mkdir()
    monkeypatch.setattr(task_history, "_task",
                        lambda root, base: Path(root))
    snapshot = artifact_source.snapshot_artifact_source(ref, task, BASE)
    assert snapshot["kind"] == "workbook" and len(snapshot["observed"]["sheets"]) == 2
    audit = skills.build_steps(_manifest("workbook_audit_report"), snapshot)
    assert len(audit[0]["data"]["rows"]) == 16
    assert any(row[3] == "公式" and row[5].startswith("=") for row in audit[0]["data"]["rows"])
    assert skills._prepared_snapshot(task, snapshot, BASE) == snapshot
    provenance_path = task / "sources" / "provenance.json"
    provenance = json.loads(provenance_path.read_text(encoding="utf-8"))
    assert provenance["kind"] == "verified_edited_workbook" and provenance["edited"] == ref
    assert len(provenance["edits"]) == 1 and provenance["root"]["sha256"]
    provenance["root"]["sha256"] = "0" * 64
    provenance_path.write_text(json.dumps(provenance), encoding="utf-8")
    with pytest.raises(ValueError, match="edited_workbook_snapshot_changed"):
        skills._prepared_snapshot(task, snapshot, BASE)


def test_development_relevance_requires_complete_audit_request():
    workbook = _manifest("workbook_audit_report")
    word = _manifest("document_structure_report")
    reference = {"version": 1, "artifact": "C:/managed/edited-spreadsheet.xlsx"}
    request = ("对所选 Excel 工作簿的全部工作表和所有单元格公式给出重算核验值，"
               "制作 Excel 及 Word 静态审计报告，不要求复制原布局")
    effect = ("全部工作表和单元格转存 Excel 与 Word 静态审计报告；"
              "公式表达式和重算核验值分列，不保留可执行公式")
    development._relevance(request, workbook, effect, reference)
    with pytest.raises(development.DevelopmentError, match="maintainer_adapter_required"):
        development._relevance("读取蓝牙设备温度并报告", workbook, effect, reference)
    with pytest.raises(development.DevelopmentError, match="development_effect_mismatch"):
        development._relevance(request, workbook, "只制作一个摘要 Word 报告，无需任何 Excel 审计输出", reference)
    with pytest.raises(development.DevelopmentError, match="development_source_kind_mismatch"):
        development._relevance("请将 Word 全部段落和表格的合并结构转存 Excel 和 Word",
                               word, "全部段落表格合并结构转存 Excel 与 Word", reference)


def test_full_workbook_audit_without_enabled_candidate_records_gap_without_model(tmp_path, monkeypatch):
    from agent_service import skill_matching

    source = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.title = "预算"
    book.active["A1"], book.active["B1"] = 2, "=A1*2"
    book.create_sheet("空白页")
    book.save(source)
    book.close()
    monkeypatch.setattr(workbook_source_host, "native_snapshot_supervised", _native)
    instruction = ("请对所选 Excel 工作簿的全部工作表、所有单元格制作逐格静态审计："
                   "在新 Excel 中列出原公式表达式和经原生重算且独立核验的值，"
                   "再生成包含全部审计行的 Word 报告。不保留可执行公式。")

    class Client:
        ledger = None

        def __init__(self, answers=()):
            self.answers = list(answers)
            self.calls = 0

        def complete(self, **_kwargs):
            self.calls += 1
            return self.answers.pop(0)

    no_call = Client()
    result = assistant.run_task(instruction, tmp_path / "tasks", source_file=source, client=no_call)
    assert result["state"] == "unsupported" and no_call.calls == 0
    assert result["source"]["kind"] == "workbook"
    gap = json.loads((Path(result["root"]) / "capability-gap.json").read_text(encoding="utf-8"))
    assert gap["development_request"]["state"] == "queued"
    preview = gap_recovery.preview_retry(result["root"], tmp_path / "tasks")
    assert preview["eligible"] and preview["source_kind"] == "workbook"
    retry = tmp_path / "tasks" / ("task-" + uuid.uuid4().hex)
    retry.mkdir()
    copied = gap_recovery.snapshot_retry_source(
        {"task_root": result["root"], "token": preview["token"]}, retry,
        tmp_path / "tasks", instruction)
    assert copied["kind"] == "workbook"
    assert workbook_source.verify_workbook_source(retry) == copied
    (retry / "request.txt").write_text(instruction, encoding="utf-8")
    assert gap_recovery.verify_prepared_origin(retry, tmp_path / "tasks")
    (retry / "events.jsonl").write_text('{"event":"started"}\n', encoding="utf-8")
    selected = _manifest("workbook_audit_report")
    monkeypatch.setattr(skills, "_selected", lambda *_args: ({}, "active", "a" * 64, selected))
    calls = []
    def claimed(*args, **_kwargs):
        calls.append(args[11]())
        return {"ok": True}
    monkeypatch.setattr(skills, "_run_claimed", claimed)
    store = tmp_path / "skill-store"
    store.mkdir()
    assert skills.run_prepared_skill(store, selected["id"], retry, tmp_path / "tasks",
                                     snapshot=copied, expected_digest="a" * 64,
                                     publish=lambda _result: None)["ok"]
    assert calls == [copied]
    native_request = retry / "sources" / "native-request.json"
    native_request.write_text(json.dumps({"sha256": "0" * 64}), encoding="utf-8")
    with pytest.raises(ValueError, match="workbook_native_request_changed"):
        skills._prepared_snapshot(retry, copied, tmp_path / "tasks")

    candidate = {"id": "audit-complete", "sha256": "a" * 64, "status": "available",
                 "recipe": "workbook_audit_report", "name": "完整审计报告", "effects": []}
    monkeypatch.setattr(skill_matching, "discover_candidates", lambda *_args, **_kwargs:
                        {"schema": 1, "candidates": [candidate], "hints": {"matched_ids": []}})
    planned = Client([{"kind": "unsupported", "reason": "无匹配能力"}] * 2)
    second = assistant.run_task(instruction, tmp_path / "other-tasks", source_file=source, client=planned)
    assert second["state"] == "unsupported" and planned.calls == 2
    partial = Client([{"kind": "source_office", "steps": []},
                      {"kind": "unsupported", "reason": "需要完整逐格适配器"}])
    third = assistant.run_task(instruction, tmp_path / "third-tasks", source_file=source, client=partial)
    assert third["state"] == "unsupported" and partial.calls == 2
    assert not any(Path(third["root"]).glob("workflow-*"))
    old_sources = Path(result["root"]) / "sources"
    for name in ("extra-one.json", "extra-two.json", "extra-three.json"):
        (old_sources / name).write_text("{}", encoding="utf-8")
    assert gap_recovery.preview_retry(result["root"], tmp_path / "tasks")["reason"] == "gap_source_invalid"


def test_preserved_real_workbook_gap_retries_with_exact_native_request(tmp_path, monkeypatch):
    if not REAL_WORKBOOK_GAP.is_dir():
        pytest.skip("preserved workbook gap unavailable")
    preview = gap_recovery.preview_retry(REAL_WORKBOOK_GAP, BASE)
    assert preview["eligible"] and preview["source_kind"] == "workbook"
    target = tmp_path / ("task-" + uuid.uuid4().hex)
    target.mkdir()
    original_task = task_history._task
    monkeypatch.setattr(task_history, "_task", lambda root, base:
                        target if Path(root).absolute() == target.absolute()
                        else original_task(root, base))
    snapshot = gap_recovery.snapshot_retry_source(
        {"task_root": str(REAL_WORKBOOK_GAP), "token": preview["token"]},
        target, BASE, preview["instruction"])
    assert snapshot["kind"] == "workbook"
    assert workbook_source.verify_workbook_source(target) == snapshot
    request = json.loads((target / "sources" / "native-request.json").read_text(encoding="utf-8"))
    assert request == {"sha256": snapshot["sha256"]}


def test_audit_review_inventory_matches_every_fixed_row():
    candidate = _manifest("workbook_audit_report")
    source = skills._fixture(candidate)
    steps = skills.build_steps(candidate, source)
    coverage = assistant._audit_review_coverage(candidate["recipe"], candidate, source, steps)
    assert [(item["name"], item["state"], item["physical_cells"])
            for item in coverage["actual_sheets"]] == [("甲", "visible", 2), ("乙", "hidden", 2)]
    assert coverage["actual_hidden_sheet_count"] == 1
    assert coverage["physical_source_cells"] == coverage["audit_rows"] == 4
    omitted = deepcopy(steps)
    omitted[0]["data"]["rows"].pop()
    with pytest.raises(assistant.AssistantError, match="skill_audit_coverage_changed"):
        assistant._audit_review_coverage(candidate["recipe"], candidate, source, omitted)
    changed = deepcopy(steps)
    changed[0]["data"]["rows"][0][4] = 2.0  # Python equality hides this type change.
    with pytest.raises(assistant.AssistantError, match="skill_audit_coverage_changed"):
        assistant._audit_review_coverage(candidate["recipe"], candidate, source, changed)

    document = _manifest("document_structure_report")
    structure = skills._fixture(document)
    structure["coverage"] = {"scope": "verified_edited_document_structure"}
    built = skills.build_steps(document, structure)
    proof = assistant._audit_review_coverage(document["recipe"], document, structure, built)
    assert proof["record_count"] == proof["audit_rows"] == 3
    assert proof["positions_in_source_order"] == [
        item["position"] for item in structure["observed"]["records"]]
    built[0]["data"]["rows"][1][0] = "body/99"
    with pytest.raises(assistant.AssistantError, match="skill_audit_coverage_changed"):
        assistant._audit_review_coverage(document["recipe"], document, structure, built)


def test_audit_intent_review_receives_only_reverified_actual_sheet_inventory(tmp_path, monkeypatch):
    from agent_service import skill_matching

    source = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.title = "预算"
    book.active["A1"], book.active["B1"] = 2, "=A1*2"
    book.create_sheet("保留页")["A1"] = "原文"
    book.save(source)
    book.close()
    monkeypatch.setattr(workbook_source_host, "native_snapshot_supervised", _native)
    base = tmp_path / "tasks"
    candidate = _manifest("workbook_audit_report")
    digest = skills.register_candidate(tmp_path / "capabilities", candidate)["sha256"]
    listed = {"id": candidate["id"], "sha256": digest, "status": "available",
              "recipe": candidate["recipe"], "name": candidate["name"], "effects": []}
    monkeypatch.setattr(skill_matching, "discover_candidates", lambda *_args, **_kwargs:
                        {"schema": 1, "candidates": [listed], "hints": {"matched_ids": []}})
    instruction = ("请把所选 Excel 工作簿的全部工作表、所有单元格做逐格静态审计，"
                   "在 Excel 记录原公式表达式和重算核验值，再生成包含全部审计行的 Word 报告，"
                   "不保留可执行公式，也不要求复制原布局。")
    class Client:
        ledger = None
        def __init__(self):
            self.requests = []
        def complete(self, **kwargs):
            self.requests.append(kwargs)
            return ({"kind": "local_skill", "id": candidate["id"], "sha256": digest}
                    if len(self.requests) == 1 else
                    {"approved": True, "reason": "核对实际两页与完整逐格记录"})
    from agent_service import local_skills
    monkeypatch.setattr(local_skills, "run_prepared_skill", lambda *_args, **kwargs:
                        kwargs["publish"]({"ok": True, "state": "succeeded", "kind": "office"}))
    client = Client()
    result = assistant.run_task(instruction, base, source_file=source, client=client)
    assert result["state"] == "succeeded" and len(client.requests) == 2
    review = json.loads(client.requests[1]["user"])
    actual = review["source_coverage"]
    assert actual["actual_sheet_count"] == 2 and actual["actual_hidden_sheet_count"] == 0
    assert [entry["name"] for entry in actual["actual_sheets"]] == ["预算", "保留页"]
    assert actual["physical_source_cells"] == actual["audit_rows"] == 3
    assert actual["all_source_sheets_and_cells_retained"] is True
    assert "Do not invent additional source sheets" in client.requests[1]["system"]

    original_plan = skill_matching.candidate_plan
    def omit_one_row(*args):
        proposed = deepcopy(original_plan(*args))
        proposed[0]["data"]["rows"].pop()
        return proposed
    monkeypatch.setattr(skill_matching, "candidate_plan", omit_one_row)
    omitted_client = Client()
    omitted = assistant.run_task(instruction, tmp_path / "omitted-tasks",
                                 source_file=source, client=omitted_client)
    assert omitted["state"] == "failed" and len(omitted_client.requests) == 1
    assert not any(Path(omitted["root"]).glob("workflow-*"))

    def corrupt_native(*args):
        planned = original_plan(*args)
        task = next((tmp_path / "corrupt-tasks").iterdir())
        target = task / "sources" / "source-native.json"
        native = json.loads(target.read_text(encoding="utf-8"))
        native["source_sha256"] = "0" * 64
        target.write_text(json.dumps(native), encoding="utf-8")
        return planned
    monkeypatch.setattr(skill_matching, "candidate_plan", corrupt_native)
    corrupt_client = Client()
    corrupted = assistant.run_task(instruction, tmp_path / "corrupt-tasks",
                                   source_file=source, client=corrupt_client)
    assert corrupted["state"] == "failed" and len(corrupt_client.requests) == 1
    assert not any(Path(corrupted["root"]).glob("workflow-*"))
