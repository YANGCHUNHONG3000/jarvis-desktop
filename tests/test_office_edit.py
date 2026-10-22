"""Precision editing, immutable originals, preservation and durable native proof."""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import zipfile

from lxml import etree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
import pytest

from agent_service import office_edit as edit


def word_fixture(path, *, complex_target=False):
    body = ('<w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:rPr><w:b/><w:sz w:val="32"/>'
            '</w:rPr><w:t>编辑验收</w:t></w:r></w:p>'
            '<w:p><w:pPr><w:spacing w:after="240"/></w:pPr><w:r><w:rPr><w:i/><w:color w:val="006699"/>'
            '</w:rPr><w:t>预算待确认</w:t></w:r>'
            + ('<w:r><w:t>复杂格式</w:t></w:r>' if complex_target else '') + '</w:p>'
            '<w:p><w:r><w:t>保留这一段和原有样式。</w:t></w:r></w:p>'
            '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>')
    types = ('<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
             '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
             '<Default Extension="xml" ContentType="application/xml"/>'
             '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
             '</Types>')
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', types)
        z.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr('word/document.xml', f'<w:document xmlns:w="{edit.W}"><w:body>{body}</w:body></w:document>')


def sheet_fixture(path):
    book = Workbook()
    ws = book.active
    ws.title = "预算"
    ws.append(["项目", "金额", "状态"])
    ws.append(["纸张", 10, "待确认"])
    ws.append(["笔", 20, "保持"])
    ws["B4"] = "=SUM(B2:B3)"
    ws["B2"].font = Font(bold=True, color="006699")
    ws["B2"].number_format = '0.00'
    ws["B2"].fill = PatternFill('solid', fgColor="FFFF00")
    ws.column_dimensions["A"].width = 24
    other = book.create_sheet("保留页")
    other["A1"] = "此页原样保留"
    other["B2"] = 123
    book.save(path)
    book.close()


def prepared(tmp_path, kind="document"):
    path = tmp_path / ("original.docx" if kind == "document" else "original.xlsx")
    (word_fixture if kind == "document" else sheet_fixture)(path)
    root = tmp_path / "input"
    root.mkdir()
    source = edit.snapshot_edit_source(path, root)
    changes = ([{"paragraph_index": 1, "before": "预算待确认", "after": "预算已确认"}]
               if kind == "document" else [{"sheet": "预算", "cell": "B2", "before": 10, "after": 15}])
    spec = {"kind": kind, "source_id": "selected-edit", "source_sha256": source["sha256"], "changes": changes}
    return path, root, source, spec


@pytest.mark.parametrize("kind", ["document", "spreadsheet"])
def test_precise_copy_and_independent_preservation(tmp_path, kind):
    path, root, source, spec = prepared(tmp_path, kind)
    before = path.read_bytes()
    output = tmp_path / ("copy" + path.suffix)
    edit._bind(spec, root)
    edit._make_copy(spec, root / ("source" + path.suffix), output)
    observed = edit.verify_preservation(path, output, spec)
    assert path.read_bytes() == before
    assert edit.verify_edit_source(root) == source
    if kind == "spreadsheet":
        book = load_workbook(output)
        assert book["预算"]["B2"].value == 15
        assert book["预算"]["B2"].font.bold
        assert book["预算"]["B2"].number_format == '0.00'
        assert book["预算"]["B4"].value == "=SUM(B2:B3)"
        assert book["保留页"]["B2"].value == 123
        book.close()
    else:
        assert observed["paragraphs"][1]["text"] == "预算已确认"


@pytest.mark.parametrize("kind", ["document", "spreadsheet"])
def test_original_and_snapshot_changes_rejected(tmp_path, kind):
    path, root, _, spec = prepared(tmp_path, kind)
    path.write_bytes(path.read_bytes() + b"changed")
    with pytest.raises(edit.EditError, match="edit_source_changed"):
        edit._bind(spec, root)


def test_complex_word_target_refused_before_output(tmp_path):
    path = tmp_path / "original.docx"
    word_fixture(path, complex_target=True)
    root = tmp_path / "input"
    root.mkdir()
    source = edit.snapshot_edit_source(path, root)
    spec = {"kind": "document", "source_id": "selected-edit", "source_sha256": source["sha256"],
            "changes": [{"paragraph_index": 1, "before": "预算待确认复杂格式", "after": "新文本"}]}
    assert source["observed"]["paragraphs"][1]["editable"] is False
    with pytest.raises(edit.EditError, match="edit_before_mismatch"):
        edit._bind(spec, root)


@pytest.mark.parametrize("change", [
    {"sheet": "预算", "cell": "B4", "before": "=SUM(B2:B3)", "after": 30},
    {"sheet": "预算", "cell": "B2", "before": 10, "after": "=WEBSERVICE(1)"},
    {"sheet": "预算", "cell": "Z99", "before": None, "after": 3},
    {"sheet": "预算", "cell": "B2", "before": 9, "after": 3},
])
def test_invalid_cells_formulas_and_stale_before_refused(tmp_path, change):
    _, root, _, spec = prepared(tmp_path, "spreadsheet")
    spec["changes"] = [change]
    with pytest.raises((edit.EditError, ValueError)):
        edit._bind(spec, root)


def test_preservation_rejects_unapproved_text_or_style_change(tmp_path):
    path, root, _, spec = prepared(tmp_path)
    output = tmp_path / "copy.docx"
    edit._make_copy(spec, root / "source.docx", output)
    parts = edit._parts(output)
    parts["word/document.xml"] = parts["word/document.xml"].replace(b'006699', b'FF0000')
    damaged = tmp_path / "damaged.docx"
    with zipfile.ZipFile(damaged, "w") as z:
        for name, content in parts.items():
            z.writestr(name, content)
    with pytest.raises(edit.EditError, match="edit_preservation_mismatch"):
        edit.verify_preservation(path, damaged, spec)


def test_native_proof_durable_revalidation_and_tampering(tmp_path, monkeypatch):
    _, root, _, spec = prepared(tmp_path)
    output = tmp_path / "output"
    output.mkdir()
    monkeypatch.setattr(edit, "_native_readback", lambda *args: {"progid": "test-only", "changes": spec["changes"]})
    result = edit.execute_edit(spec, output, root)
    assert result["ok"] is True
    assert edit.verify_edit_result(result, root, spec) == result
    receipt = Path(result["native_receipt"])
    receipt.write_bytes(receipt.read_bytes() + b" ")
    with pytest.raises(edit.EditError, match="edit_result_changed"):
        edit.verify_edit_result(result, root, spec)


def test_native_failure_unknown_and_output_never_replayed(tmp_path, monkeypatch):
    _, root, _, spec = prepared(tmp_path)
    output = tmp_path / "output"
    output.mkdir()
    def fail(*args):
        raise RuntimeError("blocked")
    monkeypatch.setattr(edit, "_native_readback", fail)
    result = edit.execute_edit(spec, output, root)
    assert result["state"] == "unknown" and result["file_verified"] is True
    with pytest.raises(ValueError, match="edit_output_conflict"):
        edit.execute_edit_supervised(spec, output, root)


def test_cancelled_before_copy(tmp_path):
    _, root, _, spec = prepared(tmp_path)
    output = tmp_path / "output"
    output.mkdir()
    result = edit.execute_edit(spec, output, root, lambda: True)
    assert result["state"] == "cancelled"
    assert not list(output.glob("*.docx"))


def test_cancellation_after_native_verification_preserves_success(tmp_path, monkeypatch):
    _, root, _, spec = prepared(tmp_path)
    output = tmp_path / "output"
    output.mkdir()
    cancelled = [False]
    def completed(*args):
        cancelled[0] = True
        return {"progid": "test-only", "changes": spec["changes"]}
    monkeypatch.setattr(edit, "_native_readback", completed)
    result = edit.execute_edit(spec, output, root, lambda: cancelled[0])
    assert result["ok"] is True
    assert edit.verify_edit_result(result, root, spec)["ok"] is True


@pytest.mark.parametrize("after", [None, True, 0, 2.5, "保留空格 "])
def test_existing_cell_scalar_types_preserve_other_parts(tmp_path, after):
    path, root, _, spec = prepared(tmp_path, "spreadsheet")
    spec["changes"][0]["after"] = after
    output = tmp_path / "copy.xlsx"
    edit._bind(spec, root)
    edit._make_copy(spec, root / "source.xlsx", output)
    edit.verify_preservation(path, output, spec)
