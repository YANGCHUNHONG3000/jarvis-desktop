"""Bounded formula edits preserve the package and verify native results."""
from types import SimpleNamespace
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import pytest

from agent_service import office_edit as edit


def prepared(tmp_path):
    source_path = tmp_path / "original.xlsx"
    book = Workbook()
    budget = book.active
    budget.title = "预算"
    budget["B2"] = 10
    budget["B3"] = 20
    budget["B4"] = "=SUM(B2:B3)"
    budget["B2"].font = Font(bold=True)
    detail = book.create_sheet("O'Brien 数据")
    detail["B2"] = 5
    detail["D4"] = "=SUM('预算'!B2:B4)+B2"
    summary = book.create_sheet("汇总")
    summary["A1"] = "='O''Brien 数据'!D4/2"
    book.save(source_path)
    book.close()
    root = tmp_path / "source"
    root.mkdir()
    public = edit.snapshot_edit_source(source_path, root)
    spec = {"kind": "spreadsheet", "source_id": "selected-edit", "source_sha256": public["sha256"],
            "changes": [{"sheet": "预算", "cell": "B2", "before": 10, "after": 15}]}
    return source_path, root, spec


def copied(tmp_path):
    original, root, spec = prepared(tmp_path)
    output = tmp_path / "edited.xlsx"
    before = original.read_bytes()
    edit._bind(spec, root)
    edit._make_copy(spec, root / "source.xlsx", output)
    observed = edit.verify_preservation(original, output, spec)
    assert original.read_bytes() == before
    return original, output, observed, spec


def install_native(monkeypatch, observed, *, bad_result=False):
    expected, order = edit._formula_expectations(observed)
    calls = []
    sheets = {}
    for sheet in observed["sheets"]:
        ranges = {}
        for cell in sheet["cells"]:
            key = sheet["name"], cell["cell"]
            formula = cell["value"] if type(cell["value"]) is dict else None
            native_formula = formula["formula"].replace("'预算'!", "预算!") if formula else None
            value = expected[key] if formula else cell["value"]
            if bad_result and key == ("汇总", "A1"):
                value = 999
            ranges[cell["cell"]] = SimpleNamespace(
                Formula=native_formula, Value2=value,
                Calculate=lambda key=key: calls.append(key))
        sheets[sheet["name"]] = SimpleNamespace(Range=lambda coordinate, ranges=ranges: ranges[coordinate])
    owned = SimpleNamespace(Worksheets=SimpleNamespace(
        Count=len(sheets), Item=lambda name: sheets[name]),Close=lambda save: calls.append(('close',save)))
    app = SimpleNamespace(Name="simulated COM protocol", Version="test-only",
                          Workbooks=SimpleNamespace(Open=lambda *args, **kwargs: owned))
    client = SimpleNamespace(DispatchEx=lambda progid: app)
    monkeypatch.setitem(sys.modules, "pythoncom", SimpleNamespace(
        CoInitialize=lambda: None, CoUninitialize=lambda: None))
    monkeypatch.setitem(sys.modules, "win32com", SimpleNamespace(client=client))
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    return calls, order


def test_cross_sheet_formula_copy_and_independent_results(tmp_path):
    _, output, observed, _ = copied(tmp_path)
    expected, order = edit._formula_expectations(observed)
    assert expected == {("预算", "B4"): 35, ("O'Brien 数据", "D4"): 75,
                        ("汇总", "A1"): 37.5}
    assert order == [("预算", "B4"), ("O'Brien 数据", "D4"), ("汇总", "A1")]
    book = load_workbook(output)
    try:
        assert book["预算"]["B2"].font.bold
        assert book["预算"]["B4"].value == "=SUM(B2:B3)"
        assert book["O'Brien 数据"]["D4"].value == "=SUM('预算'!B2:B4)+B2"
        assert book["汇总"]["A1"].value == "='O''Brien 数据'!D4/2"
    finally:
        book.close()


def test_native_formula_recalculation_and_value_proof(tmp_path, monkeypatch):
    _, output, observed, spec = copied(tmp_path)
    calls, order = install_native(monkeypatch, observed)
    receipt = edit._native_readback(output, spec, observed, lambda: None)
    assert calls == order + [('close',False)]
    assert receipt["calculation"] == "owned_formula_ranges_in_dependency_order"
    assert receipt["stored_formula_caches"] == "unchanged_in_output_package"
    assert receipt["formula_results"][-1] == {"sheet": "汇总", "cell": "A1",
                                              "expected_result": 37.5, "native_result": 37.5}


def test_native_formula_result_mismatch_rejected(tmp_path, monkeypatch):
    _, output, observed, spec = copied(tmp_path)
    install_native(monkeypatch, observed, bad_result=True)
    with pytest.raises(edit.EditError, match="edit_application_mismatch"):
        edit._native_readback(output, spec, observed, lambda: None)


@pytest.mark.parametrize("formula", ["=WEBSERVICE(1)", "=[other.xlsx]Sheet1!B1",
                                             "=B1^2", "=A1"])
def test_unsupported_or_cyclic_formula_rejected_before_native(tmp_path, formula):
    path = tmp_path / "unsafe.xlsx"
    book = Workbook()
    book.active["A1"] = formula
    book.active["B1"] = 2
    book.save(path)
    book.close()
    root = tmp_path / "source"
    root.mkdir()
    with pytest.raises(edit.EditError, match="edit_formula_unsupported"):
        edit.snapshot_edit_source(path, root)


def test_edit_that_breaks_formula_is_rejected_before_native(tmp_path):
    path = tmp_path / "ratio.xlsx"
    book = Workbook()
    book.active["A1"] = 4
    book.active["B1"] = 2
    book.active["C1"] = "=A1/B1"
    book.save(path)
    book.close()
    root = tmp_path / "source"
    root.mkdir()
    source = edit.snapshot_edit_source(path, root)
    spec = {"kind": "spreadsheet", "source_id": "selected-edit", "source_sha256": source["sha256"],
            "changes": [{"sheet": "Sheet", "cell": "B1", "before": 2, "after": 0}]}
    output = tmp_path / "edited.xlsx"
    edit._make_copy(spec, root / "source.xlsx", output)
    with pytest.raises(edit.EditError, match="edit_formula_unsupported"):
        edit.verify_preservation(path, output, spec)
