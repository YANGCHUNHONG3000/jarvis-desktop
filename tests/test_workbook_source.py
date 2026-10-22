"""Formula results come from computation and native evidence, never stale caches."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import re
import zipfile

from openpyxl import Workbook, load_workbook
import pytest

from agent_service import workbook_source as source
from agent_service import workbook_source_host as host


def fixture(path, *, formula=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "预算 明细"
    sheet.append(["项目", "金额"])
    sheet.append(["纸", 10])
    sheet.append(["笔", 20])
    sheet.append(["合计", formula or "=SUM(B2:B3)"])
    other = book.create_sheet("汇总")
    other.append(["字段", "结果"])
    other.append(["原合计", "='预算 明细'!B4"])
    other.append(["调整后", "='预算 明细'!B4*2+5"])
    other.append(["跨页求和", "=SUM('预算 明细'!B2:B3)+B2"])
    sparse = book.create_sheet("保留页")
    sparse["A1"] = "保留空白"
    sparse["B2"] = 123
    book.create_sheet("空白页")
    data = BytesIO()
    book.save(data)
    book.close()
    with zipfile.ZipFile(data) as raw, zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as out:
        for item in raw.infolist():
            content = raw.read(item)
            if item.filename.startswith("xl/worksheets/"):
                content = re.sub(rb'(<f[^>]*>.*?</f>)<v[^>]*>.*?</v>', rb'\1<v>99999</v>', content)
            out.writestr(item, content)


def receipt(path):
    observed, _ = source._parse(path)
    return {"version": 1, "source_sha256": source._digest(path), "sheets": observed["sheets"],
            "application_verified": True, "file_unchanged": True,
            "app": {"progid": "KET.Application", "name": "unit-test", "version": "test"},
            "calculation": "owned_formula_ranges_in_dependency_order"}


def prepared(tmp_path, monkeypatch):
    path = tmp_path / "fixture.xlsx"
    fixture(path)
    root = tmp_path / "task"
    root.mkdir()
    def native(path, directory, stop, deadline):
        value = receipt(path)
        source.office_source._write_json_new(directory / "source-native.json", value)
        return value
    monkeypatch.setattr(host, "native_snapshot_supervised", native)
    return path, root


def test_cross_sheet_formulas_ignore_stale_cached_values(tmp_path):
    path = tmp_path / "fixture.xlsx"
    fixture(path)
    cached = load_workbook(path, data_only=True)
    assert cached["预算 明细"]["B4"].value == 99999
    cached.close()
    observed, order = source._parse(path)
    assert observed["sheets"][0]["values"][-1] == ["合计", 30]
    assert observed["sheets"][1]["values"] == [["字段", "结果"], ["原合计", 30], ["调整后", 65], ["跨页求和", 60]]
    assert order.index(("预算 明细", "B4")) < order.index(("汇总", "B2"))
    assert observed["sheets"][2]["values"] == [["保留空白", None], [None, 123]]
    assert observed["sheets"][3] == {"name": "空白页", "range": "A1:A1", "state": "visible", "values": [[None]], "formulas": []}


def test_snapshot_complete_and_reverify_without_native_call(tmp_path, monkeypatch):
    path, root = prepared(tmp_path, monkeypatch)
    before = path.read_bytes()
    public = source.snapshot_workbook_source(path, root)
    assert public["kind"] == "workbook" and public["coverage"]["sheet_count"] == 4
    assert public == json.loads((root / "sources/source.json").read_text(encoding="utf-8"))
    assert (root / "sources/source.xlsx").read_bytes() == before == path.read_bytes()
    assert source.validate_snapshot(public) == public
    monkeypatch.setattr(host, "native_snapshot_supervised", lambda *args: pytest.fail("native replay"))
    assert source.verify_workbook_source(root) == public
    with pytest.raises(FileExistsError):
        source.snapshot_workbook_source(path, root)


@pytest.mark.parametrize("formula", ["=AVERAGE(B2:B3)", "=WEBSERVICE(B2)", "=SUM('[other.xlsx]S'!A1)",
                                      "=SUM(B2:B999)", "=B2/0", "=1E200", "=B2^2", "=SUM(Table1[Column])",
                                      "=汇总!B2", "=SUM(B2:B3" ])
def test_unsupported_or_unsafe_formulas_rejected_before_native(tmp_path, monkeypatch, formula):
    path = tmp_path / "fixture.xlsx"
    fixture(path, formula=formula)
    root = tmp_path / "task"
    root.mkdir()
    monkeypatch.setattr(host, "native_snapshot_supervised", lambda *args: pytest.fail("opened unsafe formula"))
    with pytest.raises(ValueError):
        source.snapshot_workbook_source(path, root)


@pytest.mark.parametrize("target", ["source.xlsx", "source-native.json"])
def test_hash_bound_source_and_native_receipt_tampering(tmp_path, monkeypatch, target):
    path, root = prepared(tmp_path, monkeypatch)
    source.snapshot_workbook_source(path, root)
    record = root / "sources" / target
    record.write_bytes(record.read_bytes() + b" ")
    with pytest.raises(ValueError, match="workbook_source_changed"):
        source.verify_workbook_source(root)


def test_public_formula_or_matrix_result_tampering(tmp_path, monkeypatch):
    path, root = prepared(tmp_path, monkeypatch)
    public = source.snapshot_workbook_source(path, root)
    for field in ("matrix", "formula", "expression"):
        damaged = deepcopy(public)
        sheet = damaged["observed"]["sheets"][0]
        if field == "matrix":
            sheet["values"][-1][1] = 99999
        elif field == "formula":
            sheet["formulas"][0]["result"] = 99999
        else:
            sheet["formulas"][0]["expression"] = "=SUM(B2:B3)+1"
        with pytest.raises(ValueError):
            source.validate_snapshot(damaged)


def test_native_computed_mismatch_refused(tmp_path, monkeypatch):
    path, root = prepared(tmp_path, monkeypatch)
    def wrong(path, directory, stop, deadline):
        value = receipt(path)
        value["sheets"][0]["values"][-1][1] = 99999
        source.office_source._write_json_new(directory / "source-native.json", value)
        return value
    monkeypatch.setattr(host, "native_snapshot_supervised", wrong)
    with pytest.raises(ValueError, match="workbook_native_mismatch"):
        source.snapshot_workbook_source(path, root)
    assert not (root / "sources/source.json").exists()


@pytest.mark.parametrize("expression,expected", [("=1+2*3", 7), ("=(1+2)*3", 9), ("=-B2+5", -5),
                                               ("=SUM(B2:B3,2)/2", 16), ("=1/3", 0.333333333333333)])
def test_arithmetic_precedence_and_precision(tmp_path, expression, expected):
    path = tmp_path / "fixture.xlsx"
    fixture(path, formula=expression)
    observed, _ = source._parse(path)
    assert observed["sheets"][0]["values"][-1][1] == expected


def test_recursion_limits(tmp_path):
    path = tmp_path / "fixture.xlsx"
    fixture(path, formula="=" + "(" * 66 + "1" + ")" * 66)
    with pytest.raises(ValueError, match="depth_limit"):
        source._parse(path)


def test_cancel_before_snapshot_creates_no_source_directory(tmp_path, monkeypatch):
    path, root = prepared(tmp_path, monkeypatch)
    with pytest.raises(ValueError, match="workbook_cancelled"):
        source.snapshot_workbook_source(path, root, lambda: True)
    assert not (root / "sources").exists()


def test_hidden_formula_name_refused_before_native(tmp_path, monkeypatch):
    from openpyxl.workbook.defined_name import DefinedName
    path = tmp_path / "named.xlsx"
    book = Workbook()
    book.active["A1"] = "说明"
    book.defined_names.add(DefinedName("Auto_Open", attr_text='CALL("anything")'))
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="defined_names_unsupported"):
        source._parse(path)


def test_bad_native_structure_fails_as_value_error(tmp_path):
    path = tmp_path / "fixture.xlsx"
    fixture(path)
    observed, _ = source._parse(path)
    value = receipt(path)
    value["sheets"][0] = None
    with pytest.raises(ValueError, match="workbook_native_mismatch"):
        source._native_matches(value, observed, source._digest(path))
