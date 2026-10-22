"""Read one explicitly selected Office source into an immutable, bounded snapshot.

This module does not start Office or call a model. A caller supplies the local path;
model plans may refer only to the returned fixed source ID and observations.
"""

import csv
from decimal import Decimal, InvalidOperation
from hashlib import sha256
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
import xml.etree.ElementTree as ET
import zipfile


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_MAX_SOURCE = 20_000_000
_MAX_UNCOMPRESSED = 40_000_000
_MAX_PART = 10_000_000
_MAX_ENTRIES = 1000
_MAX_COLUMNS = 12
_MAX_ROWS = 200
_NUMBER = re.compile(r"-?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?\Z")
_CELL = re.compile(r"([A-Z]+)([1-9][0-9]*)\Z")


class SourceError(ValueError):
    def __init__(self, code):
        super().__init__(code)
        self.safe_code = code


def _linked(path):
    current = Path(path).absolute()
    while True:
        if current.is_symlink() or (hasattr(current, "is_junction") and current.is_junction()):
            return True
        if current == current.parent:
            return False
        current = current.parent


def _identity(stat):
    return (stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns)


def _write_json_new(path, value):
    data = json.dumps(value, ensure_ascii=False, sort_keys=True,
                      separators=(",", ":")).encode("utf-8")
    with path.open("xb") as file:
        file.write(data)
        file.flush()
        os.fsync(file.fileno())


def _copy_stable(source_path, target):
    if _linked(source_path) or not source_path.is_file():
        raise SourceError("source_path_invalid")
    first_path_stat = source_path.stat()
    if not 0 < first_path_stat.st_size <= _MAX_SOURCE:
        raise SourceError("source_size_invalid")
    digest = sha256()
    with source_path.open("rb") as source:
        before = os.fstat(source.fileno())
        if _identity(before) != _identity(first_path_stat):
            raise SourceError("source_changed")
        copied = 0
        with target.open("xb") as snapshot:
            while True:
                chunk = source.read(1024 * 1024)
                if not chunk:
                    break
                copied += len(chunk)
                if copied > _MAX_SOURCE:
                    raise SourceError("source_size_invalid")
                digest.update(chunk)
                snapshot.write(chunk)
            snapshot.flush()
            os.fsync(snapshot.fileno())
        source.seek(0)
        reread = sha256()
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            reread.update(chunk)
        after = os.fstat(source.fileno())
    try:
        final_path_stat = source_path.stat()
    except OSError as exc:
        raise SourceError("source_changed") from exc
    if (_identity(after) != _identity(before) or _linked(source_path)
            or _identity(final_path_stat) != _identity(before)
            or copied != before.st_size or reread.hexdigest() != digest.hexdigest()):
        raise SourceError("source_changed")
    snapshot_hash = sha256(target.read_bytes()).hexdigest()
    if snapshot_hash != digest.hexdigest():
        raise SourceError("snapshot_mismatch")
    return digest.hexdigest(), copied


def _package_parts(path):
    try:
        archive = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        raise SourceError("package_invalid") from exc
    with archive:
        infos = archive.infolist()
        if len(infos) > _MAX_ENTRIES or sum(info.file_size for info in infos) > _MAX_UNCOMPRESSED:
            raise SourceError("package_size_invalid")
        seen = set()
        parts = {}
        for info in infos:
            name = info.filename
            normalized = PurePosixPath(name)
            lowered = name.casefold()
            if (not name or "\\" in name or normalized.is_absolute() or ".." in normalized.parts
                    or (not name.endswith("/") and name != str(normalized))
                    or ":" in normalized.parts[0] or lowered in seen or info.flag_bits & 1
                    or info.file_size > _MAX_PART
                    or (info.file_size > 100_000 and info.file_size > 200 * max(1, info.compress_size))):
                raise SourceError("package_unsafe")
            seen.add(lowered)
            if any(marker in lowered for marker in ("vbaproject", "activex", "embeddings", "externallinks",
                                                    "connections.xml", "querytables/", "pivotcache/")) or lowered.endswith(".bin"):
                raise SourceError("package_active_content")
            if lowered.startswith(("word/media/", "xl/media/", "word/drawings/", "xl/drawings/")):
                raise SourceError("package_unsupported_content")
            if name.endswith("/"):
                continue
            if lowered.endswith((".xml", ".rels")):
                try:
                    data = archive.read(info)
                except Exception as exc:
                    raise SourceError("package_invalid") from exc
                upper = data.upper()
                if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
                    raise SourceError("package_unsafe_xml")
                if lowered.endswith(".rels"):
                    try:
                        relations = ET.fromstring(data)
                    except ET.ParseError as exc:
                        raise SourceError("package_invalid_xml") from exc
                    for relation in relations:
                        attrs = {key.casefold(): value for key, value in relation.attrib.items()}
                        target = attrs.get("target", "").lower()
                        if (attrs.get("targetmode", "").lower() == "external"
                                or target.startswith("//") or re.match(r"[a-z][a-z0-9+.-]*:", target)):
                            raise SourceError("package_external_link")
                parts[name] = data
        content_types = parts.get("[Content_Types].xml")
        if content_types is None or b"macroenabled" in content_types.lower():
            raise SourceError("package_active_content")
        return parts


def _xml(parts, name):
    data = parts.get(name)
    if data is None:
        raise SourceError("package_missing_part")
    try:
        return ET.fromstring(data)
    except ET.ParseError as exc:
        raise SourceError("package_invalid_xml") from exc


def _word_text(node):
    pieces = []
    containers = {f"{{{_W}}}{name}" for name in ("p", "r", "tc")}
    properties = {f"{{{_W}}}{name}" for name in ("pPr", "rPr", "tcPr")}
    glyphs = {"t": None, "tab": "\t", "br": "\n", "cr": "\n",
              "noBreakHyphen": "\u2011", "softHyphen": "\u00ad"}

    def visit(element):
        if element.tag in properties:
            return
        if element.tag in containers:
            for child in element:
                visit(child)
            return
        if not element.tag.startswith(f"{{{_W}}}"):
            raise SourceError("document_unsupported_content")
        name = element.tag[len(_W) + 2:]
        if name in {"bookmarkStart", "bookmarkEnd", "proofErr"} and not len(element) and not element.text:
            return  # WPS emits non-visible bookmark/spelling range markers.
        if name not in glyphs or len(element):
            raise SourceError("document_unsupported_content")
        if name == "br" and element.get(f"{{{_W}}}type", "textWrapping") != "textWrapping":
            raise SourceError("document_unsupported_content")
        pieces.append((element.text or "") if name == "t" else glyphs[name])

    visit(node)
    return "".join(pieces)


def _document(parts):
    if any(name.startswith("word/") and (re.match(r"word/(header|footer)[0-9]*\.xml$", name)
                                            or name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"})
           for name in parts):
        raise SourceError("document_unsupported_content")
    root = _xml(parts, "word/document.xml")
    if root.tag != f"{{{_W}}}document":
        raise SourceError("document_invalid")
    body = root.find(f"{{{_W}}}body")
    if body is None:
        raise SourceError("document_invalid")
    forbidden = {"ins", "del", "moveFrom", "moveTo", "altChunk", "drawing", "pict", "object",
                 "fldSimple", "fldChar", "instrText", "delText", "commentReference",
                 "footnoteReference", "endnoteReference", "vanish", "webHidden", "sdt",
                 "gridSpan", "vMerge", "headerReference", "footerReference", "txbxContent",
                 "specVanish", "sym", "control", "subDoc"}
    if any(not element.tag.startswith(f"{{{_W}}}")
           or element.tag.rsplit("}", 1)[-1] in forbidden for element in body.iter()):
        raise SourceError("document_unsupported_content")
    paragraphs = []
    table = None
    table_seen = False
    for child in body:
        if child.tag == f"{{{_W}}}p":
            text = _word_text(child)
            if table_seen and text.strip():
                raise SourceError("document_order_unsupported")
            if table_seen:
                continue
            if len(text) > 3000:
                raise SourceError("document_size_invalid")
            paragraphs.append(text)
        elif child.tag == f"{{{_W}}}tbl":
            if table_seen:
                raise SourceError("document_tables_unsupported")
            table_seen = True
            if any(item.tag not in {f"{{{_W}}}tblPr", f"{{{_W}}}tblGrid", f"{{{_W}}}tr"}
                   for item in child):
                raise SourceError("document_unsupported_content")
            grid = []
            for row in child.findall(f"{{{_W}}}tr"):
                if any(item.tag not in {f"{{{_W}}}trPr", f"{{{_W}}}tblPrEx", f"{{{_W}}}tc"} for item in row):
                    raise SourceError("document_unsupported_content")
                cells = []
                for cell in row.findall(f"{{{_W}}}tc"):
                    if any(item.tag not in {f"{{{_W}}}tcPr", f"{{{_W}}}p"} for item in cell):
                        raise SourceError("document_unsupported_content")
                    if cell.find(f".//{{{_W}}}tbl") is not None:
                        raise SourceError("document_tables_unsupported")
                    if len(cell.findall(f"{{{_W}}}p")) != 1:
                        raise SourceError("document_tables_unsupported")
                    value = _word_text(cell)
                    if len(value) > 2000:
                        raise SourceError("document_size_invalid")
                    cells.append(value)
                grid.append(cells)
            if not grid or not 1 <= len(grid[0]) <= _MAX_COLUMNS or len(grid) > _MAX_ROWS + 1:
                raise SourceError("document_table_invalid")
            if any(len(row) != len(grid[0]) for row in grid) or any(not value.strip() for value in grid[0]):
                raise SourceError("document_table_invalid")
            table = {"headers": grid[0], "rows": grid[1:]}
        elif child.tag != f"{{{_W}}}sectPr":
            raise SourceError("document_unsupported_content")
    if len(paragraphs) > 39 or sum(map(len, paragraphs)) > 60_000 or (not paragraphs and table is None):
        raise SourceError("document_size_invalid")
    observed = {"paragraphs": paragraphs}
    if table is not None:
        observed["table"] = table
    return observed, {"scope": "main_body_only", "paragraphs": len(paragraphs),
                      "tables": 1 if table is not None else 0}


def _sheet_raw_bounds(parts):
    sheet_parts = [name for name in parts if re.fullmatch(r"xl/worksheets/[^/]+\.xml", name)]
    if len(sheet_parts) != 1:
        raise SourceError("workbook_sheets_unsupported")
    root = _xml(parts, sheet_parts[0])
    if root.find(f".//{{{_S}}}mergeCell") is not None:
        raise SourceError("workbook_merged_unsupported")
    for element in root.iter():
        if element.tag == f"{{{_S}}}f":
            raise SourceError("workbook_formula_unsupported")
        if element.tag in {f"{{{_S}}}row", f"{{{_S}}}col"} and element.get("hidden") in {"1", "true", "True"}:
            raise SourceError("workbook_hidden_unsupported")
    max_row = max_col = 0
    addresses = set()
    for cell in root.iter(f"{{{_S}}}c"):
        match = _CELL.fullmatch(cell.get("r", ""))
        if match is None:
            raise SourceError("workbook_cell_invalid")
        address = cell.get("r")
        if address in addresses or len(addresses) >= (_MAX_ROWS + 1) * _MAX_COLUMNS:
            raise SourceError("workbook_cell_invalid")
        addresses.add(address)
        value = cell.find(f"{{{_S}}}v")
        if cell.get("t") in (None, "n") and value is not None and value.text is not None:
            try:
                numeric = Decimal(value.text)
            except InvalidOperation as exc:
                raise SourceError("workbook_numeric_invalid") from exc
            if (not numeric.is_finite() or _decimal_significant_digits(numeric) > 15
                    or numeric.adjusted() >= 15):
                raise SourceError("workbook_numeric_precision_unsupported")
        col = 0
        for letter in match.group(1):
            col = col * 26 + ord(letter) - 64
        max_col = max(max_col, col)
        max_row = max(max_row, int(match.group(2)))
    if not 1 <= max_col <= _MAX_COLUMNS or not 1 <= max_row <= _MAX_ROWS + 1:
        raise SourceError("workbook_size_invalid")
    return max_row, max_col


def _cell_value(value, *, excel=False):
    from datetime import date, datetime, time
    if value is None or type(value) in (bool, str):
        if type(value) is str and len(value) > 2000:
            raise SourceError("cell_size_invalid")
        return value
    if isinstance(value, (date, datetime, time)):
        raise SourceError("workbook_date_unsupported")
    if type(value) is int:
        number = Decimal(value)
        if _decimal_significant_digits(number) > 15 or number.adjusted() >= 15:
            raise SourceError("workbook_numeric_precision_unsupported")
        return value
    if type(value) is float and math.isfinite(value):
        number = Decimal(str(value))
        if _decimal_significant_digits(number) <= 15 and number.adjusted() < 15:
            return value
        raise SourceError("workbook_numeric_precision_unsupported")
    raise SourceError("cell_type_unsupported")


def _decimal_significant_digits(value):
    if not value.is_finite():
        return 10**9
    if value.is_zero():
        return 1
    digits = value.as_tuple().digits
    # Decimal stores coefficient digits as a tuple; strip zeros without context rounding.
    while len(digits) > 1 and digits[-1] == 0:
        digits = digits[:-1]
    return len(digits)


def _table(headers, rows, scope):
    if (not 1 <= len(headers) <= _MAX_COLUMNS or len(rows) > _MAX_ROWS
            or any(type(header) is not str or not header.strip() or len(header) > 2000 for header in headers)
            or any(len(row) != len(headers) for row in rows)):
        raise SourceError("table_invalid")
    return {"headers": headers, "rows": rows}, {"scope": scope, "rows": len(rows), "columns": len(headers)}


def _workbook(path, parts):
    max_row, max_col = _sheet_raw_bounds(parts)
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise SourceError("workbook_invalid") from exc
    try:
        if len(workbook.worksheets) != 1 or workbook.worksheets[0].sheet_state != "visible":
            raise SourceError("workbook_sheets_unsupported")
        sheet = workbook.worksheets[0]
        grid = []
        for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            values = []
            for cell in row:
                if cell.data_type in {"f", "e"} or cell.is_date:
                    raise SourceError("workbook_formula_or_date_unsupported")
                values.append(_cell_value(cell.value, excel=True))
            grid.append(values)
        headers = grid[0]
        rows = grid[1:]
        observed, coverage = _table(headers, rows, "single_visible_sheet_literal_cells")
        coverage["sheet"] = sheet.title
        return observed, coverage
    finally:
        workbook.close()


def _csv_value(value):
    if len(value) > 2000:
        raise SourceError("cell_size_invalid")
    if not _NUMBER.fullmatch(value) or re.fullmatch(r"-0(?:\.0+)?", value):
        return value
    if "." in value:
        try:
            precise = Decimal(value)
        except (InvalidOperation, OverflowError, ValueError):
            return value
        if _decimal_significant_digits(precise) > 15:
            return value
        number = float(value)
        return number if math.isfinite(number) and Decimal(str(number)) == precise else value
    if len(value.lstrip("-")) > 15:
        return value
    return int(value)


def _csv(path):
    try:
        with path.open("r", encoding="utf-8-sig", errors="strict", newline="") as file:
            rows = []
            for row in csv.reader(file, strict=True):
                if len(rows) >= _MAX_ROWS + 1 or not row or len(row) > _MAX_COLUMNS:
                    raise SourceError("csv_size_invalid")
                rows.append(row)
    except (OSError, UnicodeError, csv.Error) as exc:
        raise SourceError("csv_invalid") from exc
    if not rows:
        raise SourceError("csv_size_invalid")
    headers = rows[0]
    data = [[_csv_value(value) for value in row] for row in rows[1:]]
    return _table(headers, data, "entire_utf8_csv_first_row_headers")


def _writer_compatible(display_name, kind, observed):
    """Reject an intake that cannot fit the already verified Office writer."""
    from .office import validate_spec
    spec = {"kind": "spreadsheet" if kind == "table" else "document",
            "title": display_name}
    if kind == "table":
        spec.update(headers=observed["headers"], rows=observed["rows"], sum_columns=[])
    else:
        spec["paragraphs"] = observed["paragraphs"]
        if "table" in observed:
            spec["table"] = observed["table"]
    try:
        validate_spec(spec)
    except ValueError as exc:
        raise SourceError("source_writer_unsupported") from exc


def snapshot_source(source_path, task_root) -> dict:
    """Snapshot one user-selected file and return bounded, path-free observations."""
    source = Path(source_path).absolute()
    root = Path(task_root).absolute()
    if (_linked(root) or not root.is_dir() or _linked(source) or not source.is_file()
            or source.suffix.casefold() not in {".docx", ".xlsx", ".csv"}):
        raise SourceError("source_path_invalid")
    if source.resolve().is_relative_to(root.resolve()):
        raise SourceError("source_path_invalid")
    sources = root / "sources"
    try:
        sources.mkdir(exist_ok=False)
    except FileExistsError as exc:
        raise SourceError("source_snapshot_exists") from exc
    snapshot = sources / ("source" + source.suffix.casefold())
    digest, size = _copy_stable(source, snapshot)
    suffix = source.suffix.casefold()
    if suffix == ".csv":
        observed, coverage = _csv(snapshot)
        kind = "table"
    else:
        parts = _package_parts(snapshot)
        if suffix == ".xlsx":
            observed, coverage = _workbook(snapshot, parts)
            kind = "table"
        else:
            observed, coverage = _document(parts)
            kind = "document"
    result = {"source_id": "selected", "kind": kind, "sha256": digest,
              "bytes": size, "display_name": source.name, "coverage": coverage,
              "observed": observed}
    if len(json.dumps(observed, ensure_ascii=False, separators=(",", ":"))) > 32_000:
        raise SourceError("source_observation_oversize")
    _writer_compatible(source.name, kind, observed)
    _write_json_new(sources / "source.json", result)
    return result
