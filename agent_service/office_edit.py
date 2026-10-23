"""Precise edits of bounded OOXML copies, with independent preservation checks.

Unchanged ZIP members remain byte identical. Word edits change a single text node;
Excel edits change an existing scalar cell, preserving its style and other nodes.
Native Office only opens the finished copy read-only; it never saves the source.
"""
from copy import deepcopy
from hashlib import sha256
import json
import os
from pathlib import Path
import re
import zipfile

from lxml import etree as ET

from . import office_source
from .office import _scalar, _string, _equal_scalar

W = office_source._W
S = office_source._S
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_DIGEST = re.compile(r"[0-9a-f]{64}\Z")
_CELL = re.compile(r"[A-Z]{1,2}[1-9][0-9]{0,3}\Z")
_WORD_COORDINATES = ("table_index", "row_index", "column_index", "paragraph_index")
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


class EditError(ValueError):
    def __init__(self, code):
        super().__init__(code)
        self.safe_code = code


def _hash(path):
    if office_source._linked(path) or not Path(path).is_file() or Path(path).stat().st_size > 20_000_000:
        raise EditError("edit_path_invalid")
    return sha256(Path(path).read_bytes()).hexdigest()


def _read(path):
    if office_source._linked(path) or Path(path).stat().st_size > 1_000_000:
        raise EditError("edit_record_invalid")
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _write(path, value):
    office_source._write_json_new(Path(path), value)


def _xml(data):
    return ET.fromstring(data, parser=ET.XMLParser(resolve_entities=False, no_network=True))


def _parts(path):
    # Reuse the existing bounded package security policy, including external links,
    # macros, media, encrypted entries, DTDs and ZIP expansion checks.
    office_source._package_parts(Path(path))
    with zipfile.ZipFile(path) as archive:
        return {item.filename: archive.read(item) for item in archive.infolist()}


def _sheet_parts(parts):
    workbook = _xml(parts["xl/workbook.xml"])
    rels = _xml(parts["xl/_rels/workbook.xml.rels"])
    targets = {item.get("Id"): item.get("Target") for item in rels}
    result = {}
    for sheet in workbook.findall(f"{{{S}}}sheets/{{{S}}}sheet"):
        target = targets.get(sheet.get(f"{{{R}}}id"), "")
        name = target.lstrip("/") if target.startswith("/") else "xl/" + target
        if name not in parts or not re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", name):
            raise EditError("edit_sheet_unsupported")
        result[sheet.get("name")] = name
    if not 1 <= len(result) <= 8:
        raise EditError("edit_sheet_limit")
    return result


def _word_target(paragraph):
    runs = paragraph.findall(f"{{{W}}}r")
    if len(runs) != 1:
        return None
    run = runs[0]
    texts = run.findall(f"{{{W}}}t")
    if len(texts) != 1 or any(child.tag not in {f"{{{W}}}rPr", f"{{{W}}}t"} for child in run):
        return None
    if any(child.tag not in {f"{{{W}}}pPr", f"{{{W}}}r"} for child in paragraph):
        return None
    return texts[0]


def _word_text_nodes(paragraph):
    """Return plain direct text nodes; non-text run content remains read-only."""
    if any(child.tag not in {f"{{{W}}}pPr", f"{{{W}}}r"} for child in paragraph):
        return None
    nodes = []
    for run in paragraph.findall(f"{{{W}}}r"):
        if any(child.tag not in {f"{{{W}}}rPr", f"{{{W}}}t"} for child in run):
            return None
        nodes.extend(run.findall(f"{{{W}}}t"))
    return nodes or None


def _word_edit_target(paragraph, before, after):
    """Resolve an unambiguous whole-paragraph edit to exactly one w:t node."""
    nodes = _word_text_nodes(paragraph)
    if nodes is None or "".join(node.text or "" for node in nodes) != before:
        raise EditError("edit_word_complex")

    def candidates(values, old, new):
        found = []
        prefix = ""
        for index, value in enumerate(values):
            suffix = old[len(prefix) + len(value):]
            if (new.startswith(prefix) and new.endswith(suffix)
                    and len(new) >= len(prefix) + len(suffix)):
                replacement = new[len(prefix):len(new) - len(suffix) if suffix else len(new)]
                if replacement != value:
                    found.append((index, replacement))
            prefix += value
        return found

    values = [node.text or "" for node in nodes]
    forward = candidates(values, before, after)
    if len(forward) != 1:
        raise EditError("edit_word_ambiguous" if forward else "edit_word_complex")
    index, replacement = forward[0]
    changed = list(values)
    changed[index] = replacement
    reverse = candidates(changed, after, before)
    if len(reverse) != 1 or reverse[0][0] != index:
        raise EditError("edit_word_ambiguous")
    return nodes[index], replacement


def _word_output_space(node, replacement, observation_version):
    if observation_version == 1:
        return "preserve"  # Historical single-node writer always set this.
    return ("preserve" if replacement[:1].isspace() or replacement[-1:].isspace()
            else node.get(_XML_SPACE))


def _word_body(parts):
    """Validate the editable document scope without the creation reader's table layout limits."""
    if any(name.startswith("word/") and (re.match(r"word/(header|footer)[0-9]*\.xml$", name)
            or name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}) for name in parts):
        raise EditError("edit_word_content_unsupported")
    tree = _xml(parts["word/document.xml"])
    body = tree.find(f"{{{W}}}body")
    if tree.tag != f"{{{W}}}document" or body is None:
        raise EditError("edit_word_invalid")
    forbidden = {"ins", "del", "moveFrom", "moveTo", "altChunk", "drawing", "pict", "object",
                 "fldSimple", "fldChar", "instrText", "delText", "commentReference",
                 "footnoteReference", "endnoteReference", "vanish", "webHidden", "sdt",
                 "headerReference", "footerReference", "txbxContent", "specVanish", "sym", "control", "subDoc"}
    structural_parents = {"tbl": {"body", "tc"}, "tr": {"tbl"}, "tc": {"tr"},
                          "tblGrid": {"tbl"}, "gridCol": {"tblGrid"}, "tcPr": {"tc"}}
    for element in body.iter():
        name = element.tag.rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""
        if name in structural_parents and element.getparent().tag.rsplit("}", 1)[-1] not in structural_parents[name]:
            raise EditError("edit_word_table_invalid")
        if name in {"hMerge", "gridBefore", "gridAfter"}:
            raise EditError("edit_word_merged_table_unsupported")
        if name in {"gridSpan", "vMerge"} and (element.getparent().tag != f"{{{W}}}tcPr"
                or element.getparent().getparent().tag != f"{{{W}}}tc"):
            raise EditError("edit_word_table_invalid")
        if not isinstance(element.tag, str) or not element.tag.startswith(f"{{{W}}}") or name in forbidden:
            raise EditError("edit_word_content_unsupported")
    if any(child.tag not in {f"{{{W}}}p", f"{{{W}}}tbl", f"{{{W}}}sectPr"} for child in body):
        raise EditError("edit_word_content_unsupported")
    return body


def _word_layout(body):
    """Validate logical grids and return preorder tables with unique cell anchors.

    Physical vertical continuations are retained for native verification, but never
    become edit targets. Nested paragraph indices count direct paragraphs only.
    """
    tables = []
    paragraph_count = 0

    def visit(table, depth):
        nonlocal paragraph_count
        if depth > 4 or len(tables) >= 8:
            raise EditError("edit_word_limit")
        if any(item.tag not in {f"{{{W}}}tblPr", f"{{{W}}}tblGrid", f"{{{W}}}tr"} for item in table):
            raise EditError("edit_word_content_unsupported")
        grids = table.findall(f"{{{W}}}tblGrid")
        rows = table.findall(f"{{{W}}}tr")
        if len(grids) > 1 or not 1 <= len(rows) <= 100:
            raise EditError("edit_word_table_invalid")
        width = len(grids[0].findall(f"{{{W}}}gridCol")) if grids else None
        layout = {"element": table, "width": width, "rows": []}
        tables.append(layout)
        active = {}
        for row_index, row in enumerate(rows):
            if any(item.tag not in {f"{{{W}}}trPr", f"{{{W}}}tblPrEx", f"{{{W}}}tc"} for item in row):
                raise EditError("edit_word_content_unsupported")
            cells, column, following = [], 0, {}
            for cell in row.findall(f"{{{W}}}tc"):
                if any(item.tag not in {f"{{{W}}}tcPr", f"{{{W}}}p", f"{{{W}}}tbl"} for item in cell):
                    raise EditError("edit_word_content_unsupported")
                properties = cell.findall(f"{{{W}}}tcPr")
                spans = cell.findall(f"{{{W}}}tcPr/{{{W}}}gridSpan")
                merges = cell.findall(f"{{{W}}}tcPr/{{{W}}}vMerge")
                if len(properties) > 1 or len(spans) > 1 or len(merges) > 1:
                    raise EditError("edit_word_table_invalid")
                span_text = spans[0].get(f"{{{W}}}val", "") if spans else "1"
                if not re.fullmatch(r"[1-9][0-9]?", span_text):
                    raise EditError("edit_word_table_invalid")
                span = int(span_text)
                merge = merges[0].get(f"{{{W}}}val", "continue") if merges else None
                if merge not in {None, "continue", "restart"} or column + span > 52:
                    raise EditError("edit_word_table_invalid")
                paragraphs = cell.findall(f"{{{W}}}p")
                if not 1 <= len(paragraphs) <= 100 or cell[-1].tag != f"{{{W}}}p":
                    raise EditError("edit_word_table_invalid")
                paragraph_count += len(paragraphs)
                if paragraph_count > 1000:
                    raise EditError("edit_observation_limit")
                if merge == "continue":
                    if (active.get(column) != span or cell.findall(f"{{{W}}}tbl")
                            or any(office_source._word_text(p) for p in paragraphs)):
                        raise EditError("edit_word_table_invalid")
                if merge is not None:
                    following[column] = span
                cells.append({"element": cell, "column": column, "span": span, "merge": merge})
                for nested in cell.findall(f"{{{W}}}tbl"):
                    visit(nested, depth + 1)
                column += span
            if width is None:
                width = column
                layout["width"] = width
            if not 1 <= width <= 52 or column != width:
                raise EditError("edit_word_table_invalid")
            layout["rows"].append(cells)
            active = following
        return layout

    for table in body.findall(f"{{{W}}}tbl"):
        visit(table, 1)
    return tables


def _word_observed_item(observed, change):
    if "table_index" in change:
        return next((item for item in observed.get("table_cells", [])
                     if all(item[key] == change[key] for key in _WORD_COORDINATES)), None)
    index = change["paragraph_index"]
    paragraphs = observed["paragraphs"]
    return paragraphs[index] if index < len(paragraphs) else None


def _word_paragraph(tree, change):
    body = tree.find(f"{{{W}}}body")
    try:
        if "table_index" in change:
            table = _word_layout(body)[change["table_index"]]
            row = table["rows"][change["row_index"]]
            cell = next((item["element"] for item in row if item["column"] == change["column_index"]
                         and item["merge"] != "continue"), None)
            return cell.findall(f"{{{W}}}p")[change["paragraph_index"]]
        return body.findall(f"{{{W}}}p")[change["paragraph_index"]]
    except (IndexError, AttributeError) as exc:
        raise EditError("edit_word_target_missing") from exc


def _formula_expectations(observed):
    """Evaluate bounded formula dependencies without trusting stored OOXML caches."""
    from openpyxl.utils.cell import coordinate_to_tuple
    from . import workbook_source

    bounds, literals, formulas = {}, {}, {}
    for sheet in observed["sheets"]:
        rows = columns = 0
        for cell in sheet["cells"]:
            row, column = coordinate_to_tuple(cell["cell"])
            rows, columns = max(rows, row), max(columns, column)
            key = sheet["name"], cell["cell"]
            value = cell["value"]
            if type(value) is dict:
                formulas[key] = value["formula"]
            else:
                literals[key] = value
        bounds[sheet["name"]] = rows, columns
    if not formulas:
        return {}, []
    if len(formulas) > workbook_source.MAX_FORMULAS:
        raise EditError("edit_formula_limit")
    try:
        references = set()
        def visit(node):
            if node[0] == "reference":
                references.update(node[1])
            elif node[0] == "sum":
                for child in node[1]:
                    visit(child)
            elif node[0] == "unary":
                visit(node[2])
            elif node[0] == "binary":
                visit(node[2])
                visit(node[3])
        for key, expression in formulas.items():
            visit(workbook_source._Parser(expression, key[0], bounds).parse())
            if len(references | formulas.keys()) > workbook_source.MAX_CELLS:
                raise EditError("edit_formula_limit")
        for key in references:
            if key not in formulas:
                literals.setdefault(key, None)
        evaluator = workbook_source._Evaluator(literals, formulas, bounds)
        for key in formulas:
            evaluator.cell(key)
        return {key: evaluator.values[key] for key in formulas}, evaluator.order
    except EditError:
        raise
    except (workbook_source.WorkbookSourceError, ValueError) as exc:
        raise EditError("edit_formula_unsupported") from exc


def _observe(path, kind, observation_version=1):
    if (type(observation_version) is not int or observation_version not in (1, 2)
            or (kind != "document" and observation_version != 1)):
        raise EditError("edit_observation_version_invalid")
    parts = _parts(path)
    if kind == "document":
        body = _word_body(parts)
        paragraphs = body.findall(f"{{{W}}}p")
        tables = _word_layout(body)
        if len(paragraphs) > 100 or len(tables) > 8 or not (paragraphs or tables):
            raise EditError("edit_word_limit")
        editable = (lambda p: _word_text_nodes(p) is not None) if observation_version == 2 else (
            lambda p: _word_target(p) is not None)
        observed = {"paragraphs": [{"index": i, "text": office_source._word_text(p),
                                    "editable": editable(p)}
                                   for i, p in enumerate(paragraphs)]}
        if tables:
            cells = []
            for table_index, table in enumerate(tables):
                for row_index, row in enumerate(table["rows"]):
                    for descriptor in row:
                        if descriptor["merge"] == "continue":
                            continue
                        column_index = descriptor["column"]
                        cell_paragraphs = descriptor["element"].findall(f"{{{W}}}p")
                        for paragraph_index, paragraph in enumerate(cell_paragraphs):
                            cells.append(dict(zip(_WORD_COORDINATES, (table_index, row_index, column_index,
                                                                      paragraph_index)),
                                              text=office_source._word_text(paragraph),
                                              editable=editable(paragraph)))
                            if len(cells) > 1000:
                                raise EditError("edit_observation_limit")
            observed["table_cells"] = cells
        items = observed["paragraphs"] + observed.get("table_cells", [])
        if any(len(item["text"]) > 3000 for item in items) or sum(len(item["text"]) for item in items) > 60_000:
            raise EditError("edit_word_limit")
        if len(json.dumps(observed, ensure_ascii=False)) > 64_000:
            raise EditError("edit_observation_limit")
        return observed
    import openpyxl
    book = openpyxl.load_workbook(path, data_only=False, keep_links=False)
    sheets = _sheet_parts(parts)
    output = []
    count = 0
    try:
        for name, part in sheets.items():
            sheet = book[name]
            if sheet.max_row > 500 or sheet.max_column > 52:
                raise EditError("edit_sheet_limit")
            cells = []
            for node in _xml(parts[part]).iter(f"{{{S}}}c"):
                coordinate = node.get("r")
                cell = sheet[coordinate]
                if cell.data_type == "f":
                    value = {"formula": cell.value}
                else:
                    try:
                        value = _scalar(cell.value)
                    except ValueError as exc:
                        raise EditError("edit_cell_type_unsupported") from exc
                cells.append({"cell": coordinate, "value": value,
                              "editable": cell.data_type != "f" and not any(
                                  coordinate in merged for merged in sheet.merged_cells.ranges)})
            count += len(cells)
            if count > 2000:
                raise EditError("edit_cell_limit")
            output.append({"name": name, "cells": cells})
    finally:
        book.close()
    observed = {"sheets": output}
    if any(type(cell["value"]) is dict for sheet in output for cell in sheet["cells"]):
        workbook = _xml(parts["xl/workbook.xml"])
        if workbook.findall(f"{{{S}}}definedNames/{{{S}}}definedName"):
            raise EditError("edit_formula_unsupported")
        _formula_expectations(observed)
    if len(json.dumps(observed, ensure_ascii=False)) > 64_000:
        raise EditError("edit_observation_limit")
    return observed


def snapshot_edit_source(path, root):
    source = Path(path).absolute()
    root = Path(root).absolute()
    if (office_source._linked(root) or not root.is_dir() or source.suffix.lower() not in {".docx", ".xlsx"}
            or source.resolve().is_relative_to(root.resolve())):
        raise EditError("edit_source_invalid")
    suffix = source.suffix.lower()
    snapshot = root / ("source" + suffix)
    digest, size = office_source._copy_stable(source, snapshot)
    kind = "document" if suffix == ".docx" else "spreadsheet"
    observation_version = 2 if kind == "document" else 1
    result = {"source_id": "selected-edit", "kind": kind, "sha256": digest,
              "bytes": size, "display_name": source.name,
              "observed": _observe(snapshot, kind, observation_version)}
    if kind == "document":
        result["observation_version"] = observation_version
    _write(root / "source.json", {"public": result, "original": str(source)})
    return result


def verify_edit_source(source_root):
    root = Path(source_root)
    record = _read(root / "source.json")
    public = record["public"]
    suffix = ".docx" if public["kind"] == "document" else ".xlsx"
    path = root / ("source" + suffix)
    observation_version = public.get("observation_version", 1)
    if (type(observation_version) is not int or observation_version not in (1, 2)
            or (public["kind"] != "document" and observation_version != 1)):
        raise EditError("edit_source_changed")
    if (_hash(path) != public["sha256"] or _hash(record["original"]) != public["sha256"]
            or _observe(path, public["kind"], observation_version) != public["observed"]):
        raise EditError("edit_source_changed")
    return deepcopy(public)


def validate_edit_spec(spec):
    if (type(spec) is not dict or set(spec) != {"kind", "source_id", "source_sha256", "changes"}
            or spec["kind"] not in {"document", "spreadsheet"} or spec["source_id"] != "selected-edit"
            or type(spec["source_sha256"]) is not str or not _DIGEST.fullmatch(spec["source_sha256"])
            or type(spec["changes"]) is not list or not 1 <= len(spec["changes"]) <= 20):
        raise EditError("edit_spec_invalid")
    result = deepcopy(spec)
    seen = set()
    for change in result["changes"]:
        if type(change) is not dict:
            raise EditError("edit_change_invalid")
        if spec["kind"] == "document":
            coordinates = _WORD_COORDINATES if "table_index" in change else ("paragraph_index",)
            limits = {"table_index": 8, "row_index": 100, "column_index": 52, "paragraph_index": 100}
            if (set(change) != set(coordinates) | {"before", "after"}
                    or any(type(change[key]) is not int or not 0 <= change[key] < limits[key] for key in coordinates)):
                raise EditError("edit_change_invalid")
            key = tuple(change[name] for name in coordinates)
            for field in ("before", "after"):
                _string(change[field], 3000, "edit_text")
            if not change["after"]:
                raise EditError("edit_empty_word_unsupported")
        else:
            if (set(change) != {"sheet", "cell", "before", "after"}
                    or type(change["sheet"]) is not str or type(change["cell"]) is not str
                    or not _CELL.fullmatch(change["cell"])):
                raise EditError("edit_change_invalid")
            key = (change["sheet"], change["cell"])
            for field in ("before", "after"):
                change[field] = _scalar(change[field])
                if isinstance(change[field], str) and change[field].startswith("="):
                    raise EditError("edit_formula_target_forbidden")
            if change["after"] == "":
                raise EditError("edit_clear_requires_null")
        if key in seen or (type(change["before"]) is type(change["after"]) and change["before"] == change["after"]):
            raise EditError("edit_duplicate_or_noop")
        seen.add(key)
    return result


def _bind(spec, source_root):
    spec = validate_edit_spec(spec)
    source = verify_edit_source(source_root)
    if source["kind"] != spec["kind"] or source["sha256"] != spec["source_sha256"]:
        raise EditError("edit_source_mismatch")
    tree = (_xml(_parts(Path(source_root) / "source.docx")["word/document.xml"])
            if spec["kind"] == "document" and source.get("observation_version") == 2 else None)
    for change in spec["changes"]:
        if spec["kind"] == "document":
            item = _word_observed_item(source["observed"], change)
            value = item["text"] if item else None
        else:
            item = next((cell for sheet in source["observed"]["sheets"] if sheet["name"] == change["sheet"]
                         for cell in sheet["cells"] if cell["cell"] == change["cell"]), None)
            value = item["value"] if item else None
        if item is None or item["editable"] is not True or not _equal_scalar(value, change["before"]):
            raise EditError("edit_before_mismatch")
        if tree is not None:
            _word_edit_target(_word_paragraph(tree, change), change["before"], change["after"])
    return spec, source


def _masked(parts, spec, word_indices=None):
    result = dict(parts)
    grouped = {}
    if spec["kind"] == "document":
        grouped["word/document.xml"] = spec["changes"]
    else:
        sheets = _sheet_parts(parts)
        for change in spec["changes"]:
            grouped.setdefault(sheets[change["sheet"]], []).append(change)
    for part, changes in grouped.items():
        tree = _xml(parts[part])
        for position, change in enumerate(changes):
            if spec["kind"] == "document":
                paragraph = _word_paragraph(tree, change)
                if word_indices is None:
                    current = office_source._word_text(paragraph)
                    other = change["after"] if current == change["before"] else change["before"]
                    node, _ = _word_edit_target(paragraph, current, other)
                else:
                    nodes = _word_text_nodes(paragraph)
                    if nodes is None or word_indices[position] >= len(nodes):
                        raise EditError("edit_preservation_mismatch")
                    node = nodes[word_indices[position]]
                node.text = "__EDIT_VALUE__"
                node.attrib.pop(_XML_SPACE, None)
            else:
                matches = tree.findall(f".//{{{S}}}c[@r='{change['cell']}']")
                if len(matches) != 1:
                    raise EditError("edit_cell_missing")
                node = matches[0]
                node.attrib.pop("t", None)
                for child in list(node):
                    if child.tag in {f"{{{S}}}v", f"{{{S}}}is"}:
                        node.remove(child)
                    elif child.tag == f"{{{S}}}f":
                        raise EditError("edit_formula_target_forbidden")
        result[part] = ET.tostring(tree, method="c14n")
    return result


def verify_preservation(source_path, artifact, spec, observation_version=1):
    """Independent check of every ZIP member and every unapproved XML node."""
    spec = validate_edit_spec(spec)
    before, after = _parts(source_path), _parts(artifact)
    word_indices = None
    if spec["kind"] == "document":
        source_tree = _xml(before["word/document.xml"])
        output_tree = _xml(after["word/document.xml"])
        word_indices = []
        for change in spec["changes"]:
            source_paragraph = _word_paragraph(source_tree, change)
            if office_source._word_text(source_paragraph) != change["before"]:
                raise EditError("edit_before_mismatch")
            output_paragraph = _word_paragraph(output_tree, change)
            node, replacement = _word_edit_target(source_paragraph, change["before"], change["after"])
            index = _word_text_nodes(source_paragraph).index(node)
            output_nodes = _word_text_nodes(output_paragraph)
            if (output_nodes is None or index >= len(output_nodes)
                    or (output_nodes[index].text or "") != replacement
                    or office_source._word_text(output_paragraph) != change["after"]
                    or output_nodes[index].get(_XML_SPACE) !=
                    _word_output_space(node, replacement, observation_version)):
                raise EditError("edit_preservation_mismatch")
            word_indices.append(index)
    if _masked(before, spec, word_indices) != _masked(after, spec, word_indices):
        raise EditError("edit_preservation_mismatch")
    observed = _observe(artifact, spec["kind"], observation_version)
    for change in spec["changes"]:
        if spec["kind"] == "document":
            item = _word_observed_item(observed, change)
            if item is None:
                raise EditError("edit_word_target_missing")
            value = item["text"]
        else:
            value = next(cell["value"] for sheet in observed["sheets"] if sheet["name"] == change["sheet"]
                         for cell in sheet["cells"] if cell["cell"] == change["cell"])
        if not _equal_scalar(value, change["after"]):
            raise EditError("edit_file_value_mismatch")
    return observed


def _make_copy(spec, source, target, observation_version=1):
    parts = _parts(source)
    modified = {}
    if spec["kind"] == "document":
        tree = _xml(parts["word/document.xml"])
        for change in spec["changes"]:
            node, replacement = _word_edit_target(_word_paragraph(tree, change),
                                                  change["before"], change["after"])
            space = _word_output_space(node, replacement, observation_version)
            node.text = replacement
            if space is None:
                node.attrib.pop(_XML_SPACE, None)
            else:
                node.set(_XML_SPACE, space)
        modified["word/document.xml"] = ET.tostring(tree, encoding="UTF-8", xml_declaration=True)
    else:
        import openpyxl
        # Let the mature reader resolve shared strings and validate scalar types;
        # patch just the existing cell to avoid library package normalization.
        book = openpyxl.load_workbook(source, data_only=False, keep_links=False)
        sheets = _sheet_parts(parts)
        trees = {}
        try:
            for change in spec["changes"]:
                part = sheets[change["sheet"]]
                tree = trees.setdefault(part, _xml(parts[part]))
                node = tree.find(f".//{{{S}}}c[@r='{change['cell']}']")
                cell = book[change["sheet"]][change["cell"]]
                cell.value = change["after"]
                for child in list(node):
                    if child.tag in {f"{{{S}}}v", f"{{{S}}}is"}:
                        node.remove(child)
                node.attrib.pop("t", None)
                value = cell.value
                if isinstance(value, str):
                    node.set("t", "inlineStr")
                    text = ET.SubElement(ET.SubElement(node, f"{{{S}}}is"), f"{{{S}}}t")
                    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    text.text = value
                elif value is not None:
                    if type(value) is bool:
                        node.set("t", "b")
                    ET.SubElement(node, f"{{{S}}}v").text = str(int(value)) if type(value) is bool else str(value)
            modified = {part: ET.tostring(tree, encoding="UTF-8", xml_declaration=True)
                        for part, tree in trees.items()}
        finally:
            book.close()
    with zipfile.ZipFile(source) as original, Path(target).open("xb") as output:
        with zipfile.ZipFile(output, "w") as copied:
            copied.comment = original.comment
            for info in original.infolist():
                copied.writestr(info, modified.get(info.filename, parts[info.filename]))
        output.flush()
        os.fsync(output.fileno())


def _word_native_semantics(body):
    """Compare native document meaning without depending on run/style normalization."""
    layouts = {item["element"]: item for item in _word_layout(body)}

    def content(parent):
        return tuple(("p", office_source._word_text(child)) if child.tag == f"{{{W}}}p"
                     else table(child) for child in parent if child.tag in {f"{{{W}}}p", f"{{{W}}}tbl"})

    def table(element):
        layout = layouts[element]
        # WPS removes grid divisions which are covered by a span in EVERY row.
        # Preserve all actual cell edges and merge topology while allowing that
        # native normalization. Source targets still use the original grid.
        edges = sorted({0, layout["width"]} | {edge for row in layout["rows"] for cell in row
                                              for edge in (cell["column"], cell["column"] + cell["span"])})
        positions = {edge: index for index, edge in enumerate(edges)}
        return ("table", len(edges)-1, tuple(tuple(
            (positions[cell["column"]], positions[cell["column"] + cell["span"]] - positions[cell["column"]],
             cell["merge"], content(cell["element"])) for cell in row)
            for row in layout["rows"]))

    return content(body)


def _word_native_stream(body):
    # WPS/Word includes a row terminator as well as each physical cell terminator.
    # Vertical continuations remain physical empty cells in this native surface.
    def content(parent, in_cell=False):
        blocks = [child for child in parent if child.tag in {f"{{{W}}}p", f"{{{W}}}tbl"}]
        result = []
        for index, child in enumerate(blocks):
            if child.tag == f"{{{W}}}p":
                result.append(office_source._word_text(child) + ("\r\x07" if in_cell and index == len(blocks)-1 else "\r"))
            else:
                for row in child.findall(f"{{{W}}}tr"):
                    result.extend(content(cell, True) for cell in row.findall(f"{{{W}}}tc"))
                    result.append("\r\x07")
        return "".join(result)
    return content(body)


def _native_complex_word(owned, body, check):
    check()
    native_xml = str(owned.Content.WordOpenXML)
    if len(native_xml) > 4_000_000 or "<!DOCTYPE" in native_xml or "<!ENTITY" in native_xml:
        raise EditError("edit_application_mismatch")
    package = _xml(native_xml.encode("utf-8"))
    pkg = "{http://schemas.microsoft.com/office/2006/xmlPackage}"
    documents = package.findall(f"{pkg}part[@{pkg}name='/word/document.xml']/{pkg}xmlData/{{{W}}}document")
    if len(documents) != 1:
        raise EditError("edit_application_mismatch")
    native_body = documents[0].find(f"{{{W}}}body")
    if native_body is None or _word_native_semantics(native_body) != _word_native_semantics(body):
        raise EditError("edit_application_mismatch")
    check()
    if str(owned.Content.Text).replace("\v", "\n") != _word_native_stream(body):
        raise EditError("edit_application_mismatch")


def _native_readback(path, spec, observed, check):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    owned = None
    try:
        check()
        progid = "KWPS.Application" if spec["kind"] == "document" else "KET.Application"
        app = win32com.client.DispatchEx(progid)
        app.Visible = True
        check()
        owned = (app.Documents.Open(str(path), ReadOnly=True) if spec["kind"] == "document"
                 else app.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True))
        check()
        values = []
        if spec["kind"] == "document":
            body = _word_body(_parts(path))
            complex_table = (any(element.tag in {f"{{{W}}}gridSpan", f"{{{W}}}vMerge"} for element in body.iter())
                             or any(table.find(f".//{{{W}}}tbl") is not None for table in body.findall(f"{{{W}}}tbl")))
            if complex_table:
                _native_complex_word(owned, body, check)
            else:
                _native_simple_word(owned, observed, check)
            values = [{**{key: c[key] for key in (_WORD_COORDINATES if "table_index" in c else ("paragraph_index",))},
                       "after": c["after"]} for c in spec["changes"]]
        else:
            from . import workbook_source
            expected_results, order = _formula_expectations(observed)
            if int(owned.Worksheets.Count) != len(observed["sheets"]):
                raise EditError("edit_application_mismatch")
            for name, coordinate in order:
                check()
                owned.Worksheets.Item(name).Range(coordinate).Calculate()
            formula_results = []
            bounds = {}
            from openpyxl.utils.cell import coordinate_to_tuple
            for sheet in observed["sheets"]:
                rows = columns = 0
                for cell in sheet["cells"]:
                    row, column = coordinate_to_tuple(cell["cell"])
                    rows, columns = max(rows, row), max(columns, column)
                bounds[sheet["name"]] = rows, columns
            for sheet in observed["sheets"]:
                native = owned.Worksheets.Item(sheet["name"])
                for cell in sheet["cells"]:
                    check()
                    actual = native.Range(cell["cell"])
                    expected = cell["value"]
                    if type(expected) is dict:
                        key = sheet["name"], cell["cell"]
                        value = actual.Value2
                        if (not workbook_source._same_native_formula(str(actual.Formula), expected["formula"],
                                                                     sheet["name"], bounds)
                                or not _equal_scalar(value, expected_results[key])):
                            raise EditError("edit_application_mismatch")
                        formula_results.append({"sheet": key[0], "cell": key[1],
                                                "expected_result": expected_results[key],
                                                "native_result": value})
                    elif not _equal_scalar(actual.Value2, expected):
                        raise EditError("edit_application_mismatch")
            values = [{"sheet": c["sheet"], "cell": c["cell"], "after": c["after"]} for c in spec["changes"]]
        result = {"progid": progid, "name": str(app.Name), "version": str(app.Version), "changes": values}
        if spec["kind"] == "spreadsheet" and formula_results:
            result.update(formula_results=formula_results,
                          calculation="owned_formula_ranges_in_dependency_order",
                          stored_formula_caches="unchanged_in_output_package")
        return result
    finally:
        try:
            if owned is not None:
                owned.Close(False)
        finally:
            pythoncom.CoUninitialize()


def _native_simple_word(owned, observed, check):
    # Word's document-wide Paragraphs includes cell paragraphs. Body
    # indices deliberately count only direct body paragraphs in OOXML.
    body_paragraphs = []
    for index in range(1, owned.Paragraphs.Count + 1):
        check()
        paragraph_range = owned.Paragraphs.Item(index).Range
        if not paragraph_range.Information(12):  # wdWithInTable
            body_paragraphs.append(str(paragraph_range.Text).replace("\v", "\n"))
    if body_paragraphs != [paragraph["text"] + "\r" for paragraph in observed["paragraphs"]]:
        raise EditError("edit_application_mismatch")
    table_cells = observed.get("table_cells", [])
    table_count = max((item["table_index"] for item in table_cells), default=-1) + 1
    if owned.Tables.Count != table_count:
        raise EditError("edit_application_mismatch")
    for table_index in range(table_count):
        table = owned.Tables.Item(table_index + 1)
        expected_cells = [item for item in table_cells if item["table_index"] == table_index]
        row_count = max(item["row_index"] for item in expected_cells) + 1
        column_count = max(item["column_index"] for item in expected_cells) + 1
        if table.Rows.Count != row_count or table.Columns.Count != column_count:
            raise EditError("edit_application_mismatch")
        for row_index in range(row_count):
            for column_index in range(column_count):
                check()
                paragraphs = table.Cell(row_index + 1, column_index + 1).Range.Paragraphs
                expected = [item for item in expected_cells if item["row_index"] == row_index
                            and item["column_index"] == column_index]
                if paragraphs.Count != len(expected):
                    raise EditError("edit_application_mismatch")
                for item in expected:
                    check()
                    index = item["paragraph_index"]
                    text = str(paragraphs.Item(index + 1).Range.Text).replace("\v", "\n")
                    ending = "\r\x07" if index == len(expected) - 1 else "\r"
                    if text != item["text"] + ending:
                        raise EditError("edit_application_mismatch")


def execute_edit(spec, root, source_root, stop=lambda: False):
    result = {"ok": False, "state": "failed", "file_verified": False,
              "application_verified": False, "preservation_verified": False}
    stage = "validate"
    native = False
    def mark(value):
        _write(Path(root) / ("phase-" + value + ".json"), {"stage": value})
    def check():
        if stop():
            raise EditError("edit_cancelled")
    try:
        mark(stage)
        spec, source = _bind(spec, source_root)
        root = Path(root).resolve(strict=True)
        suffix = ".docx" if spec["kind"] == "document" else ".xlsx"
        artifact = root / ("edited-document.docx" if suffix == ".docx" else "edited-spreadsheet.xlsx")
        original = Path(source_root) / ("source" + suffix)
        check()
        stage = "copy_edit"
        mark(stage)
        _make_copy(spec, original, artifact, source.get("observation_version", 1))
        check()
        observed = verify_preservation(original, artifact, spec, source.get("observation_version", 1))
        result.update(artifact=str(artifact), path=str(artifact), sha256=_hash(artifact),
                      source_sha256=source["sha256"], file_verified=True, preservation_verified=True)
        stage = "native_readback"
        mark(stage)
        native = True
        app = _native_readback(artifact, spec, observed, check)
        # Once all native values were read successfully, later cancellation
        # cannot erase that evidence. Finish the read-only file/source checks.
        verify_edit_source(source_root)
        verify_preservation(original, artifact, spec, source.get("observation_version", 1))
        result.update(ok=True, state="succeeded", application_verified=True, original_unchanged=True,
                      app=app, observed=observed)
        receipt = {"spec": spec, "source_sha256": source["sha256"], "artifact": str(artifact),
                   "sha256": result["sha256"], "app": app, "observed": observed}
        receipt_path = root / "native-edit-receipt.json"
        _write(receipt_path, receipt)
        result.update(native_receipt=str(receipt_path), native_receipt_sha256=_hash(receipt_path))
    except Exception as exc:
        result.update(ok=False, state="unknown" if native else "cancelled" if str(exc) == "edit_cancelled" else "failed",
                      code=str(exc) if isinstance(exc, (EditError, office_source.SourceError)) else "edit_operation_failed",
                      stage=stage)
    return result


def verify_edit_result(result, source_root, spec):
    spec, source = _bind(spec, source_root)
    if (type(result) is not dict or result.get("ok") is not True or result.get("state") != "succeeded"
            or any(result.get(key) is not True for key in ("file_verified", "application_verified",
                                                         "preservation_verified", "original_unchanged"))):
        raise EditError("edit_result_invalid")
    artifact = Path(result["artifact"])
    expected_name = "edited-document.docx" if spec["kind"] == "document" else "edited-spreadsheet.xlsx"
    receipt_path = artifact.parent / "native-edit-receipt.json"
    if (artifact.name != expected_name or result["native_receipt"] != str(receipt_path)
            or _hash(artifact) != result["sha256"] or _hash(receipt_path) != result["native_receipt_sha256"]):
        raise EditError("edit_result_changed")
    receipt = _read(receipt_path)
    if (receipt != {"spec": spec, "source_sha256": source["sha256"], "artifact": str(artifact),
                   "sha256": result["sha256"], "app": result["app"], "observed": result["observed"]}
            or result["source_sha256"] != source["sha256"]):
        raise EditError("edit_receipt_mismatch")
    suffix = ".docx" if spec["kind"] == "document" else ".xlsx"
    observed = verify_preservation(Path(source_root) / ("source" + suffix), artifact, spec,
                                   source.get("observation_version", 1))
    if observed != result["observed"]:
        raise EditError("edit_result_changed")
    return deepcopy(result)


def execute_edit_supervised(spec, root, source_root, stop=lambda: False, deadline=None):
    from .office_edit_host import execute_edit_supervised as supervised
    return supervised(spec, root, source_root, stop, deadline)
