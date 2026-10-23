"""Complete bounded workbook snapshots with independently checked native results.

Formula caches are never inputs. Only SUM, arithmetic and bounded local workbook
references are understood; unsupported semantics fail before native Office opens.
"""
from copy import deepcopy
from decimal import Decimal, InvalidOperation, localcontext
from hashlib import sha256
import json
import math
from pathlib import Path
import re

from . import office_source
from .office import _equal_scalar, _scalar, _string
from .office_edit import _parts, _sheet_parts, _xml

S = office_source._S
_REF = re.compile(r"(?:(?:'((?:[^']|'')+)'|([^'!]+))!)?(\$?[A-Z]{1,2}\$?[1-9][0-9]{0,3})(?::(\$?[A-Z]{1,2}\$?[1-9][0-9]{0,3}))?\Z", re.I)
_HEX = re.compile(r"[0-9a-f]{64}\Z")
MAX_CELLS = 1600
MAX_FORMULAS = 200


class WorkbookSourceError(ValueError):
    def __init__(self, code):
        super().__init__(code)
        self.safe_code = code


def _fail(code):
    raise WorkbookSourceError(code)


def _digest(path):
    path = Path(path)
    if office_source._linked(path) or not path.is_file() or not 0 < path.stat().st_size <= 20_000_000:
        _fail("workbook_path_invalid")
    return sha256(path.read_bytes()).hexdigest()


def _read(path):
    path = Path(path)
    if office_source._linked(path) or path.stat().st_size > 500_000:
        _fail("workbook_record_invalid")
    return json.loads(path.read_text(encoding="utf-8"))


def _number(value):
    if value is None:
        return Decimal(0)
    if isinstance(value, Decimal):
        result = value
    elif type(value) in (int, float, bool):
        result = Decimal(int(value)) if type(value) is bool else Decimal(str(value))
    else:
        _fail("workbook_formula_operand_unsupported")
    if not result.is_finite() or (not result.is_zero() and abs(result.adjusted()) > 100):
        _fail("workbook_formula_number_limit")
    return result


def _result(value):
    if value is None:
        return 0  # Excel's direct reference to an empty scalar cell is zero.
    if isinstance(value, list):
        _fail("workbook_formula_array_unsupported")
    if not isinstance(value, Decimal):
        return _scalar(value)
    value = _number(value)
    with localcontext() as context:
        context.prec = 220
        if value:
            value = value.quantize(Decimal(1).scaleb(value.adjusted() - 14))
    if value == value.to_integral() and len(str(abs(int(value)))) <= 15:
        return int(value)
    return _scalar(float(value))


class _Parser:
    def __init__(self, expression, sheet, bounds):
        from openpyxl.formula import Tokenizer
        if type(expression) is not str or not expression.startswith("=") or len(expression) > 1000:
            _fail("workbook_formula_unsupported")
        try:
            self.tokens = [token for token in Tokenizer(expression).items if token.type != "WHITE-SPACE"]
        except Exception as exc:
            raise WorkbookSourceError("workbook_formula_unsupported") from exc
        if not 1 <= len(self.tokens) <= 256:
            _fail("workbook_formula_token_limit")
        self.index, self.depth = 0, 0
        self.sheet, self.bounds = sheet, bounds

    def peek(self, value=None):
        if self.index >= len(self.tokens):
            return False
        return self.tokens[self.index] if value is None else self.tokens[self.index].value == value

    def take(self, value=None):
        token = self.peek()
        if not token or (value is not None and token.value != value):
            _fail("workbook_formula_unsupported")
        self.index += 1
        return token

    def parse(self):
        result = self.expression()
        if self.peek():
            _fail("workbook_formula_unsupported")
        return result

    def expression(self, minimum=0):
        self.depth += 1
        if self.depth > 64:
            _fail("workbook_formula_depth_limit")
        token = self.take()
        if token.type == "OPERATOR-PREFIX" and token.value in {"+", "-"}:
            node = ("unary", token.value, self.expression(3))
        elif token.type == "PAREN" and token.subtype == "OPEN":
            node = self.expression()
            self.take(")")
        elif token.type == "FUNC" and token.subtype == "OPEN" and token.value.upper() == "SUM(":
            args = []
            if not self.peek(")"):
                args.append(self.expression())
                while self.peek(","):
                    self.take(",")
                    args.append(self.expression())
            self.take(")")
            node = ("sum", args)
        elif token.type == "OPERAND" and token.subtype == "NUMBER":
            if len(token.value) > 50:
                _fail("workbook_formula_number_limit")
            try:
                number = _number(Decimal(token.value))
            except InvalidOperation as exc:
                raise WorkbookSourceError("workbook_formula_unsupported") from exc
            node = ("number", number)
        elif token.type == "OPERAND" and token.subtype == "RANGE":
            node = ("reference", self.reference(token.value))
        else:
            _fail("workbook_formula_unsupported")
        precedence = {"+": 1, "-": 1, "*": 2, "/": 2}
        while self.peek() and self.peek().type == "OPERATOR-INFIX":
            op = self.peek().value
            if op not in precedence:
                _fail("workbook_formula_unsupported")
            rank = precedence[op]
            if rank < minimum:
                break
            self.take()
            node = ("binary", op, node, self.expression(rank + 1))
        self.depth -= 1
        return node

    def reference(self, value):
        from openpyxl.utils.cell import range_boundaries, get_column_letter
        match = _REF.fullmatch(value)
        if match is None:
            _fail("workbook_formula_reference_unsupported")
        quoted, plain, first, last = match.groups()
        name = quoted.replace("''", "'") if quoted is not None else plain if plain is not None else self.sheet
        canonical = {key.casefold(): key for key in self.bounds}.get(name.casefold())
        if canonical is None:
            _fail("workbook_formula_sheet_missing")
        first, last = first.replace("$", "").upper(), (last or first).replace("$", "").upper()
        c1, r1, c2, r2 = range_boundaries(first + ":" + last)
        rows, columns = self.bounds[canonical]
        if r1 > r2 or c1 > c2 or r2 > rows or c2 > columns:
            _fail("workbook_formula_range_outside_snapshot")
        return [(canonical, f"{get_column_letter(column)}{row}")
                for row in range(r1, r2 + 1) for column in range(c1, c2 + 1)]


class _Evaluator:
    def __init__(self, literals, formulas, bounds):
        self.values, self.formulas = dict(literals), formulas
        self.ast = {key: _Parser(expression, key[0], bounds).parse() for key, expression in formulas.items()}
        self.active, self.order, self.steps = set(), [], 0

    def tick(self):
        self.steps += 1
        if self.steps > 100_000:
            _fail("workbook_formula_step_limit")

    def cell(self, key):
        self.tick()
        if key in self.values:
            return self.values[key]
        if key in self.active:
            _fail("workbook_formula_cycle")
        if len(self.active) >= 64:
            _fail("workbook_formula_depth_limit")
        self.active.add(key)
        try:
            with localcontext() as context:
                context.prec = 40
                value = _result(self.node(self.ast[key], 0))
            self.values[key] = value
            self.order.append(key)
            return value
        finally:
            self.active.remove(key)

    def node(self, node, depth):
        self.tick()
        if depth > 64:
            _fail("workbook_formula_depth_limit")
        kind = node[0]
        if kind == "number":
            return node[1]
        if kind == "reference":
            values = [self.cell(key) for key in node[1]]
            return values[0] if len(values) == 1 else values
        if kind == "sum":
            result = Decimal(0)
            for child in node[1]:
                value = self.node(child, depth + 1)
                for scalar in value if isinstance(value, list) else [value]:
                    self.tick()
                    if isinstance(scalar, Decimal) or type(scalar) in (int, float):
                        result = _number(result + _number(scalar))
            return result
        if kind == "unary":
            value = _number(self.node(node[2], depth + 1))
            return value if node[1] == "+" else -value
        left, right = (_number(self.node(child, depth + 1)) for child in node[2:])
        op = node[1]
        if op == "+":
            return _number(left + right)
        if op == "-":
            return _number(left - right)
        if op == "*":
            return _number(left * right)
        if not right:
            _fail("workbook_formula_division_by_zero")
        return _number(left / right)


def _analyze(sheets, *, verify_results=False):
    from openpyxl.utils.cell import get_column_letter, range_boundaries
    if type(sheets) is not list or not 1 <= len(sheets) <= 8:
        _fail("workbook_sheet_limit")
    bounds, literals, formulas, supplied, count = {}, {}, {}, {}, 0
    for sheet in sheets:
        if type(sheet) is not dict or set(sheet) != {"name", "range", "state", "values", "formulas"}:
            _fail("workbook_sheet_invalid")
        name = _string(sheet["name"], 31, "sheet_name", nonempty=True)
        if (name.casefold() in {key.casefold() for key in bounds} or type(sheet["state"]) is not str
                or sheet["state"] not in {"visible", "hidden", "veryHidden"}):
            _fail("workbook_sheet_invalid")
        grid = sheet["values"]
        if (type(grid) is not list or not 1 <= len(grid) <= 201 or type(grid[0]) is not list
                or not 1 <= len(grid[0]) <= 12 or any(type(row) is not list or len(row) != len(grid[0]) for row in grid)):
            _fail("workbook_grid_limit")
        rows, columns = len(grid), len(grid[0])
        if sheet["range"] != f"A1:{get_column_letter(columns)}{rows}":
            _fail("workbook_range_invalid")
        count += rows * columns
        if count > MAX_CELLS:
            _fail("workbook_cell_limit")
        bounds[name] = rows, columns
        for r, row in enumerate(grid, 1):
            for c, value in enumerate(row, 1):
                literals[name, f"{get_column_letter(c)}{r}"] = _scalar(value)
        if type(sheet["formulas"]) is not list:
            _fail("workbook_formulas_invalid")
        for formula in sheet["formulas"]:
            if type(formula) is not dict or set(formula) != {"cell", "expression", "result"}:
                _fail("workbook_formulas_invalid")
            cell = formula["cell"]
            if type(cell) is not str or not re.fullmatch(r"[A-Z]{1,2}[1-9][0-9]{0,3}", cell):
                _fail("workbook_formula_cell_invalid")
            key = name, cell
            if key not in literals or key in formulas:
                _fail("workbook_formula_cell_invalid")
            formulas[key] = formula["expression"]
            supplied[key] = formula["result"]
    if len(formulas) > MAX_FORMULAS:
        _fail("workbook_formula_limit")
    for key in formulas:
        literals.pop(key)
    evaluator = _Evaluator(literals, formulas, bounds)
    for key in formulas:
        evaluator.cell(key)
    result = deepcopy(sheets)
    for sheet in result:
        for r, row in enumerate(sheet["values"], 1):
            for c in range(1, len(row) + 1):
                key = sheet["name"], f"{get_column_letter(c)}{r}"
                value = evaluator.values[key]
                if verify_results and not _equal_scalar(row[c - 1], value):
                    _fail("workbook_formula_result_mismatch")
                row[c - 1] = value
        for formula in sheet["formulas"]:
            key = sheet["name"], formula["cell"]
            if verify_results and not _equal_scalar(supplied[key], evaluator.values[key]):
                _fail("workbook_formula_result_mismatch")
            formula["result"] = evaluator.values[key]
    if len(json.dumps(result, ensure_ascii=False, separators=(",", ":"))) > 64_000:
        _fail("workbook_observation_limit")
    return result, evaluator.order


def _parse(path):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
    parts = _parts(path)
    workbook_xml = _xml(parts["xl/workbook.xml"])
    if workbook_xml.findall(f"{{{S}}}definedNames/{{{S}}}definedName"):
        # Names may contain executable legacy formulas or external expressions.
        # This bounded reader has no named-formula semantics, so reject them
        # before Office opens rather than silently trusting hidden definitions.
        _fail("workbook_defined_names_unsupported")
    sheet_parts = _sheet_parts(parts)
    count = 0
    # Inspect the real cells before constructing the mature library's workbook.
    for part in sheet_parts.values():
        seen = set()
        for node in _xml(parts[part]).iter(f"{{{S}}}c"):
            coordinate = node.get("r", "")
            if not re.fullmatch(r"[A-Z]{1,2}[1-9][0-9]{0,3}", coordinate) or coordinate in seen:
                _fail("workbook_cell_invalid")
            row, column = coordinate_to_tuple(coordinate)
            if row > 201 or column > 12:
                _fail("workbook_grid_limit")
            seen.add(coordinate)
            count += 1
            if count > MAX_CELLS:
                _fail("workbook_cell_limit")
    book = load_workbook(path, data_only=False, keep_links=False)
    sheets = []
    try:
        if set(sheet_parts) != set(book.sheetnames) or len(book.worksheets) != len(book.sheetnames):
            _fail("workbook_sheet_unsupported")
        for sheet in book.worksheets:
            if not 1 <= sheet.max_row <= 201 or not 1 <= sheet.max_column <= 12:
                _fail("workbook_grid_limit")
            if sheet.merged_cells.ranges:
                _fail("workbook_merged_cells_unsupported")
            rows, columns = sheet.max_row, sheet.max_column
            grid, formulas = [], []
            for cells in sheet.iter_rows(min_row=1, max_row=rows, min_col=1, max_col=columns):
                values = []
                for cell in cells:
                    if cell.data_type == "f":
                        formulas.append({"cell": cell.coordinate, "expression": cell.value, "result": None})
                        values.append(None)  # Stored <v> is deliberately ignored.
                    elif cell.data_type == "e" or cell.is_date:
                        _fail("workbook_cell_type_unsupported")
                    else:
                        values.append(_scalar(cell.value))
                grid.append(values)
            sheets.append({"name": sheet.title, "range": f"A1:{get_column_letter(columns)}{rows}",
                           "state": sheet.sheet_state, "values": grid, "formulas": formulas})
    finally:
        book.close()
    observed, order = _analyze(sheets)
    return {"sheets": observed}, order


def validate_snapshot(public):
    if (type(public) is not dict or set(public) != {"source_id", "kind", "sha256", "bytes", "display_name",
                                                    "coverage", "observed", "native_sha256"}
            or public["source_id"] != "selected" or public["kind"] != "workbook"
            or type(public["bytes"]) is not int or not 0 < public["bytes"] <= 20_000_000
            or any(type(public[key]) is not str or not _HEX.fullmatch(public[key]) for key in ("sha256", "native_sha256"))
            or type(public["observed"]) is not dict or set(public["observed"]) != {"sheets"}):
        _fail("workbook_snapshot_invalid")
    _string(public["display_name"], 255, "display_name", nonempty=True)
    sheets, _ = _analyze(public["observed"]["sheets"], verify_results=True)
    if public["coverage"] != _coverage(sheets):
        _fail("workbook_coverage_invalid")
    return deepcopy(public)


def _coverage(sheets):
    return {"scope": "all_sheets_all_rows", "sheet_count": len(sheets),
            "formula_results": "native_recalculated_and_independently_verified",
            "numeric_precision": "15_significant_digits"}


def _same_native_formula(actual, expected, sheet, bounds):
    """Allow native sheet-name quoting normalization, not merely equal results."""
    if actual == expected:
        return True
    def identity(expression):
        parser = _Parser(expression, sheet, bounds)
        parser.parse()  # Both sides must still obey the bounded formula grammar.
        result = []
        for token in parser.tokens:
            value = token.value
            if token.type == "OPERAND" and token.subtype == "RANGE":
                match = _REF.fullmatch(value)
                quoted, plain, first, last = match.groups()
                name = quoted.replace("''", "'") if quoted is not None else plain
                # Preserve relative/absolute anchors and local-vs-qualified form.
                value = (name.casefold() if name is not None else None,
                         first.upper(), last.upper() if last is not None else None)
            result.append((token.type, token.subtype, value))
        return result
    try:
        return identity(actual) == identity(expected)
    except (WorkbookSourceError, ValueError, TypeError, AttributeError):
        return False


def _native_matches(receipt, observed, digest):
    if (type(receipt) is not dict or type(receipt.get("version")) is not int or receipt.get("version") != 1 or receipt.get("source_sha256") != digest
            or receipt.get("application_verified") is not True or receipt.get("file_unchanged") is not True
            or type(receipt.get("app")) is not dict or receipt["app"].get("progid") != "KET.Application"
            or type(receipt.get("sheets")) is not list or len(receipt["sheets"]) != len(observed["sheets"])):
        _fail("workbook_native_receipt_invalid")
    bounds = {sheet["name"]: (len(sheet["values"]), len(sheet["values"][0]))
              for sheet in observed["sheets"]}
    for actual, expected in zip(receipt["sheets"], observed["sheets"]):
        if type(actual) is not dict or any(actual.get(key) != expected[key] for key in ("name", "range", "state")):
            _fail("workbook_native_mismatch")
        values = actual.get("values")
        if type(values) is not list or len(values) != len(expected["values"]):
            _fail("workbook_native_mismatch")
        for row, target in zip(values, expected["values"]):
            if type(row) is not list or len(row) != len(target) or any(
                    not _equal_scalar(value, desired) for value, desired in zip(row, target)):
                _fail("workbook_native_mismatch")
        formulas = actual.get("formulas")
        if type(formulas) is not list or len(formulas) != len(expected["formulas"]):
            _fail("workbook_native_mismatch")
        for formula, wanted in zip(formulas, expected["formulas"]):
            if (type(formula) is not dict or formula.get("cell") != wanted["cell"]
                    or not _same_native_formula(formula.get("expression"), wanted["expression"], expected["name"], bounds)
                    or not _equal_scalar(formula.get("result"), wanted["result"])):
                _fail("workbook_native_mismatch")


def _native_snapshot(path, stop=lambda: False):
    import pythoncom
    import win32com.client
    observed, order = _parse(path)
    initial = _digest(path)
    def check():
        if stop():
            _fail("workbook_cancelled")
    check()
    pythoncom.CoInitialize()
    owned = None
    try:
        app = win32com.client.DispatchEx("KET.Application")
        check()
        owned = app.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, ReadOnly=True)
        check()
        if int(owned.Worksheets.Count) != len(observed["sheets"]):
            _fail("workbook_native_mismatch")
        for name, cell in order:
            check()
            owned.Worksheets.Item(name).Range(cell).Calculate()
        result = []
        for expected in observed["sheets"]:
            check()
            sheet = owned.Worksheets.Item(expected["name"])
            state = {-1: "visible", 0: "hidden", 2: "veryHidden"}.get(int(sheet.Visible))
            values = sheet.Range(expected["range"]).Value2
            row_count, column_count = len(expected["values"]), len(expected["values"][0])
            if row_count == column_count == 1 and not isinstance(values, (list, tuple)):
                values = [[values]]
            else:
                values = [list(row) for row in values]
            formulas = []
            for formula in expected["formulas"]:
                check()
                cell = sheet.Range(formula["cell"])
                formulas.append({"cell": formula["cell"], "expression": str(cell.Formula), "result": cell.Value2})
            result.append({"name": str(sheet.Name), "range": expected["range"], "state": state,
                           "values": values, "formulas": formulas})
        receipt = {"version": 1, "source_sha256": initial, "sheets": result,
                   "application_verified": True, "file_unchanged": _digest(path) == initial,
                   "app": {"progid": "KET.Application", "name": str(app.Name), "version": str(app.Version)},
                   "calculation": "owned_formula_ranges_in_dependency_order"}
        _native_matches(receipt, observed, initial)
        return receipt
    finally:
        if owned is not None:
            owned.Close(SaveChanges=False)
        # Never quit a shared WPS server or touch another workbook.
        pythoncom.CoUninitialize()


def snapshot_workbook_source(path, task_root, stop=lambda: False, deadline=None):
    from .workbook_source_host import native_snapshot_supervised
    source, root = Path(path).absolute(), Path(task_root).absolute()
    if (office_source._linked(root) or not root.is_dir() or source.suffix.casefold() != ".xlsx"
            or source.resolve().is_relative_to(root.resolve())):
        _fail("workbook_source_invalid")
    if stop():
        _fail("workbook_cancelled")
    directory = root / "sources"
    directory.mkdir(exist_ok=False)
    snapshot = directory / "source.xlsx"
    digest, size = office_source._copy_stable(source, snapshot)
    observed, _ = _parse(snapshot)
    receipt = native_snapshot_supervised(snapshot, directory, stop, deadline)
    if _digest(source) != digest or _digest(snapshot) != digest:
        _fail("workbook_source_changed")
    _native_matches(receipt, observed, digest)
    public = {"source_id": "selected", "kind": "workbook", "sha256": digest,
              "bytes": size, "display_name": source.name, "coverage": _coverage(observed["sheets"]),
              "observed": observed, "native_sha256": _digest(directory / "source-native.json")}
    validate_snapshot(public)
    office_source._write_json_new(directory / "source.json", public)
    return public


def verify_workbook_source(task_root):
    directory = Path(task_root) / "sources"
    public = validate_snapshot(_read(directory / "source.json"))
    request_path = directory / "native-request.json"
    if request_path.exists() or office_source._linked(request_path):
        request = _read(request_path)
        if type(request) is not dict or request != {"sha256": public["sha256"]}:
            _fail("workbook_native_request_changed")
    if (_digest(directory / "source.xlsx") != public["sha256"]
            or (directory / "source.xlsx").stat().st_size != public["bytes"]
            or _digest(directory / "source-native.json") != public["native_sha256"]):
        _fail("workbook_source_changed")
    observed, _ = _parse(directory / "source.xlsx")
    if observed != public["observed"]:
        _fail("workbook_source_changed")
    _native_matches(_read(directory / "source-native.json"), observed, public["sha256"])
    return public
